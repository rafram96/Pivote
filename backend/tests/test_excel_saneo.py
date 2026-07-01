"""El parche de openpyxl (importado vía scripts.generar_excel) cierra los DOS
vectores de crash del Excel, sin red:
  1. caracteres de control (\\x00-\\x1f) de texto OCR/PDF → IllegalCharacterError.
  2. list/dict en una celda (el espejo permite `requisitos` como record(any)) →
     ValueError "Cannot convert [...] to Excel".
Ambos tumbaban el entregable ENTERO; ahora se sanean al vuelo."""
import openpyxl

import scripts.generar_excel  # noqa: F401 — importar aplica el parche global


def test_celda_con_control_chars_no_revienta():
    wb = openpyxl.Workbook()
    wb.active.cell(1, 1, "texto\x00con\x07control")
    assert wb.active.cell(1, 1).value == "textoconcontrol"


def test_celda_con_lista_se_coerce_a_string():
    wb = openpyxl.Workbook()
    wb.active.cell(1, 1, ["Ingeniero Civil", "Arquitecto"])
    assert wb.active.cell(1, 1).value == "Ingeniero Civil; Arquitecto"


def test_celda_con_dict_se_coerce_a_string():
    wb = openpyxl.Workbook()
    wb.active.cell(1, 1, {"cargos_validos": "A", "tipo_obra": "B"})
    assert "cargos_validos: A" in wb.active.cell(1, 1).value


def test_escalares_intactos():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, 42)
    ws.cell(1, 2, 3.14)
    ws.cell(1, 3, "texto limpio")
    ws.cell(1, 4, None)
    assert ws.cell(1, 1).value == 42
    assert ws.cell(1, 2).value == 3.14
    assert ws.cell(1, 3).value == "texto limpio"
    assert ws.cell(1, 4).value is None
