"""
Tests offline de los entregables: Excel final (CLAUDE + Base de Datos + hojas
por profesional con hitos) y ZIP InfoObras (árbol de 4 niveles).
"""
from __future__ import annotations

import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from entregables import construir_zip_infoobras, generar_excel_final, inventariar
from scripts.generar_excel import generar_excel
from schemas.cargo import etiqueta_hoja
from reglas import dias_efectivos_profesional

RAIZ = Path(__file__).resolve().parents[2]
FIXTURES = RAIZ / "fixtures"

# Espejo sintético: 2 profesionales; el prof 1 reusa el caso golden JEFE
# (certificado 2021-05-13→2023-05-06 con 3 ventanas de paralización).
ESPEJO = {
    "_meta": {"analisis_id": "t-entregables", "concurso": "CP-TEST/2026",
              "postor": "POSTOR DEMO"},
    "postor": {},
    "profesionales": [
        {"n_prof": 1, "cargo": "JEFE DE SUPERVISIÓN", "nombre": "Profesional Uno",
         "colegiatura": "CIP 11111",
         "experiencias": [
             {"n": 1, "proyecto": "Obra A", "cui": "2418877",
              "fecha_inicial": "2021-05-13", "fecha_final": "2023-05-06",
              "dias": 724, "incluye_covid": "NO", "folio": "100"},
             {"n": 2, "proyecto": "Obra B",
              "fecha_inicial": "2023-11-22", "fecha_final": "2025-07-22",
              "dias": 609, "incluye_covid": "NO", "folio": "120"},
         ]},
        {"n_prof": 2, "cargo": "ESP. ESTRUCTURAS", "nombre": "Profesional Dos",
         "experiencias": [
             {"n": 1, "proyecto": "Obra C",
              "fecha_inicial": "POR VERIFICAR (ilegible)", "fecha_final": "2020-01-01"},
         ]},
    ],
    "resumen_evaluacion": {"factores": [], "puntaje_total": None},
}

PARALIZACIONES = {
    (1, 1): [(date(2021, 10, 1), date(2021, 10, 31)),
             (date(2022, 2, 1), date(2022, 4, 30)),
             (date(2022, 7, 1), date(2022, 9, 30))],
}


# ── Excel final ───────────────────────────────────────────────────────────────

def test_excel_final_estructura_de_hojas(tmp_path):
    salida = generar_excel_final(ESPEJO, tmp_path / "final.xlsx", PARALIZACIONES)
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames[0] == "CLAUDE"
    assert wb.sheetnames[1] == "Base de Datos"
    # Pestañas cortas: prefijo genérico fuera, cada palabra a la mitad, Title Case.
    # "Jefe" se conserva (jerarquía); "ESP." no (todos son especialistas).
    assert wb.sheetnames[2] == "P1. Jefe Superv"
    assert wb.sheetnames[3] == "P2. Estruc"


def test_excel_final_base_datos_con_filtro_y_colores(tmp_path):
    salida = generar_excel_final(ESPEJO, tmp_path / "final.xlsx", PARALIZACIONES)
    ws = openpyxl.load_workbook(salida)["Base de Datos"]
    assert ws.cell(1, 1).value == "CARGO AL QUE POSTULA"
    # 3 experiencias = 3 filas de datos
    assert ws.cell(2, 1).value == "JEFE DE SUPERVISIÓN"
    assert ws.cell(4, 1).value == "ESP. ESTRUCTURAS"
    assert ws.cell(5, 1).value is None
    assert ws.auto_filter.ref is not None        # con filtros
    # color de fondo distinto por profesional
    assert ws.cell(2, 1).fill.fgColor.rgb != ws.cell(4, 1).fill.fgColor.rgb


def test_excel_final_hoja_profesional_cuadro_de_hitos(tmp_path):
    salida = generar_excel_final(ESPEJO, tmp_path / "final.xlsx", PARALIZACIONES)
    ws = openpyxl.load_workbook(salida)[etiqueta_hoja(1, "JEFE DE SUPERVISIÓN")]
    celdas = [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]
    texto = "\n".join(celdas)

    # el marco del certificado existe, con periodo certificado, paralizaciones y tramos
    assert "CERT. N°1" in texto and "DATOS DE LA EXPERIENCIA" in texto
    assert "Periodo certificado" in texto
    assert "Paralización 1 de la obra (InfoObras)" in texto
    assert "Tramo efectivo 1" in texto
    # los 4 tramos del golden JEFE: 141 + 92 + 61 + 218
    assert "EFECTIVO EXPERIENCIA 1" in texto
    valores = [c.value for fila in ws.iter_rows() for c in fila]
    for esperado in (141, 92, 61, 218, 512):
        assert esperado in valores, f"falta el valor {esperado} del golden"
    # resumen Paso 5 reconcilia con el motor de reglas
    res = dias_efectivos_profesional(
        [(date(2021, 5, 13), date(2023, 5, 6)), (date(2023, 11, 22), date(2025, 7, 22))],
        {0: PARALIZACIONES[(1, 1)]},
    )
    assert res.dias_efectivos in valores
    assert "DÍAS EFECTIVOS" in texto


def test_excel_final_muestra_periodos_sin_valorizacion(tmp_path):
    # un periodo con tipo 'sin_valorizacion' (hueco) debe verse etiquetado como
    # "obra parada" en la hoja del profesional, no como "Paralización"
    paral = {(1, 1): [
        {"inicio": date(2021, 10, 1), "fin": date(2021, 10, 31), "tipo": "paralizado"},
        {"inicio": date(2022, 1, 1), "fin": date(2022, 3, 31), "tipo": "sin_valorizacion"},
    ]}
    salida = generar_excel_final(ESPEJO, tmp_path / "final.xlsx", paral)
    ws = openpyxl.load_workbook(salida)[etiqueta_hoja(1, "JEFE DE SUPERVISIÓN")]
    texto = "\n".join(str(c.value) for f in ws.iter_rows() for c in f if c.value)
    assert "Paralización 1 de la obra (InfoObras)" in texto
    assert "Sin valorización 1 — obra parada (InfoObras)" in texto


def test_excel_final_valorizaciones_resalta_solo_meses_del_certificado(tmp_path):
    # datos reales de la captura del ingeniero: Hospital Ernesto Guzman (CUI
    # 2198319, código 102951, finalizado). Certificado 01/05/19–31/10/19 → solo
    # las valorizaciones de may–oct 2019 deben quedar en amarillo.
    espejo = {
        "_meta": {"analisis_id": "v", "concurso": "CP-V/2026", "postor": "P"},
        "postor": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "JEFE", "nombre": "N",
            "experiencias": [{"n": 1, "proyecto": "Hospital Ernesto Guzman", "folio": "600",
                              "fecha_inicial": "2019-05-01", "fecha_final": "2019-10-31"}],
        }],
        "resumen_evaluacion": {"factores": []},
    }
    fichas = {(1, 1): {
        "codigo_infoobras": "102951", "cui": "2198319", "estado": "Finalizado",
        "monto": 11131864.36, "fecha_inicio": "2019-03-25", "fecha_fin": "2019-07-23",
        "valorizaciones": [
            {"anio": 2020, "mes": 2, "estado": "En ejecución", "fisico_real": 1.0, "valorizado_real": 11131864.36},
            {"anio": 2019, "mes": 10, "estado": "En ejecución", "fisico_real": 0.5172, "valorizado_real": 4974255.19},
            {"anio": 2019, "mes": 5, "estado": "En ejecución", "fisico_real": 0.1564, "valorizado_real": 1738809.21},
            {"anio": 2019, "mes": 3, "estado": "En ejecución", "fisico_real": 0.0216, "valorizado_real": 240053.31},
        ],
    }}
    salida = generar_excel_final(espejo, tmp_path / "v.xlsx", fichas=fichas)
    ws = openpyxl.load_workbook(salida)[etiqueta_hoja(1, "JEFE")]

    color = {}
    for fila in ws.iter_rows():
        for c in fila:
            if c.value and " / " in str(c.value):
                color[str(c.value)] = (c.fill.fgColor.rgb or "")

    # may y oct 2019 → dentro del certificado → amarillo
    assert color["2019 / MAYO"].endswith("FFFF00")
    assert color["2019 / OCTUBRE"].endswith("FFFF00")
    # mar 2019 y feb 2020 → fuera → sin resaltar
    assert not color["2019 / MARZO"].endswith("FFFF00")
    assert not color["2020 / FEBRERO"].endswith("FFFF00")

    # la ficha de la obra (código, CUI, estado) aparece en la hoja
    texto = "\n".join(str(c.value) for f in ws.iter_rows() for c in f if c.value)
    assert "102951" in texto and "2198319" in texto and "Finalizado" in texto
    assert "VALORIZACIONES" in texto


def test_excel_final_fechas_no_computables_quedan_anotadas(tmp_path):
    salida = generar_excel_final(ESPEJO, tmp_path / "final.xlsx", PARALIZACIONES)
    ws = openpyxl.load_workbook(salida)[etiqueta_hoja(2, "ESP. ESTRUCTURAS")]
    texto = "\n".join(str(c.value) for fila in ws.iter_rows() for c in fila if c.value)
    assert "no computables" in texto


def test_excel_final_fixture_libertador(tmp_path):
    ruta = FIXTURES / "new_format" / "libertador_espejo.json"
    if not ruta.exists():
        pytest.skip("fixture local no disponible")
    import json
    espejo = json.loads(ruta.read_text(encoding="utf-8"))
    salida = generar_excel_final(espejo, tmp_path / "libertador_final.xlsx")
    wb = openpyxl.load_workbook(salida)
    # CLAUDE + Base de Datos + 14 hojas de profesional
    assert len(wb.sheetnames) == 2 + 14
    ws = wb["Base de Datos"]
    filas_datos = sum(1 for f in ws.iter_rows(min_row=2) if f[0].value is not None)
    assert filas_datos == 41  # 41 experiencias consolidadas


# ── ZIP InfoObras ─────────────────────────────────────────────────────────────

def test_zip_arbol_de_4_niveles_con_y_sin_documentos(tmp_path):
    # docs reales descargados por el prototipo (gitignored) si están; si no, sintéticos
    origen = RAIZ / "tools" / "_descargas_72056"
    if not origen.is_dir():
        origen = tmp_path / "docs"
        (origen / "Cronograma").mkdir(parents=True)
        (origen / "Cronograma" / "doc1.pdf").write_bytes(b"%PDF-1.4 demo")

    salida = construir_zip_infoobras(
        ESPEJO, {(1, 1): origen}, tmp_path / "infoobras.zip")

    with zipfile.ZipFile(salida) as zf:
        nombres = zf.namelist()
        # carpeta corta "P{nn} {cargo}" / "E{n}" (nombres cortos para el límite 260
        # de Windows; el nivel del concurso se omite, va en el indice.txt)
        pref = "P01 SUPERVISIÓN/E1/"
        assert any(n.startswith(pref) for n in nombres)
        # la experiencia CON descargas trae archivos reales
        con_docs = [n for n in nombres if n.startswith(pref) and not n.endswith("/")]
        assert len(con_docs) >= 1
        # las experiencias SIN descargas llevan la nota, no desaparecen
        sin_docs = [n for n in nombres if n.endswith("SIN_DOCUMENTOS.txt")]
        assert len(sin_docs) == 2  # exp (1,2) y (2,1)
        contenido = zf.read(sin_docs[0]).decode("utf-8")
        assert "Motivo" in contenido
        # índice en la raíz del ZIP
        assert "indice.txt" in nombres
        indice = zf.read("indice.txt").decode("utf-8")
        assert "SIN DOCUMENTOS" in indice


# ── Inventario (parseo offline del HTML de InfoObras) ───────────────────────

HTML_INVENTARIO = """
<html><script>
var lAvances = [
  {"lImgValorizacion": [{"UrlImg": "Doc/documento1.pdf", "nombreArchivo": "valorizacion enero",
    "Extension": "pdf", "EsFisico": 0}],
   "lImgFisico": [{"UrlImg": "Img/foto1.jpg", "nombreArchivo": "avance fisico",
    "Extension": "jpg", "EsFisico": 1}]}
];
</script>
<body>
<button data-download-url="/InfobrasWeb/Mapa/DownloadFile?filename=expediente%2Fexp.pdf&amp;name=expediente%20tecnico&amp;extension=.pdf"></button>
</body></html>
"""


def test_inventariar_distingue_documentos_de_imagenes():
    inv = inventariar(HTML_INVENTARIO)
    docs = {d["filename"] for d in inv["documentos"]}
    imgs = {i["filename"] for i in inv["imagenes"]}
    assert "Doc/documento1.pdf" in docs
    assert "expediente/exp.pdf" in docs       # botón data-download-url
    assert "Img/foto1.jpg" in imgs            # imagen excluida de documentos
    # una imagen NUNCA se cuela como documento aunque fuera .pdf
    assert not docs & imgs


HTML_POR_AVANCE = """
<html><script>
var lAvances = [
  {"Anio": "2025", "Mes": "FEBRERO",
   "lImgValorizacion": [{"UrlImg": "doc/val12.pdf", "nombreArchivo": "VALORIZACION DE OBRA Nº 0012",
     "Extension": "pdf", "EsFisico": 0}], "lImgFisico": []},
  {"Anio": "2025", "Mes": "ENERO",
   "lImgValorizacion": [{"UrlImg": "doc/val11.pdf", "nombreArchivo": "VALORIZACION Nº 0011",
     "Extension": "pdf", "EsFisico": 0}], "lImgFisico": []}
];
</script>
<body>
<button data-download-url="/InfobrasWeb/Mapa/DownloadFile?filename=expediente%2Fet.pdf&amp;name=expediente&amp;extension=.pdf"></button>
</body></html>
"""


def test_inventariar_por_avance_agrupa_por_hito():
    from entregables.zip_infoobras import inventariar_por_avance, _etiqueta_hito
    inv = inventariar_por_avance(HTML_POR_AVANCE)
    assert len(inv["avances"]) == 2
    feb = next(a for a in inv["avances"] if a["mes"] == "FEBRERO")
    assert feb["anio"] == 2025
    assert feb["documentos"][0]["nombre"] == "VALORIZACION DE OBRA Nº 0012"
    # carpeta del hito, ordenable cronológicamente
    assert _etiqueta_hito(2025, "FEBRERO") == "2025-02 FEBRERO"
    # el expediente es obra-level (no por avance)
    assert any(d["nombre"] == "expediente" for d in inv["obra"]["documentos"])
    assert all(not a_doc["filename"].startswith("expediente")
               for a in inv["avances"] for a_doc in a["documentos"])


def test_descargar_a_carpeta_prefija_nombre_con_fecha(tmp_path):
    """El documento de valorización se guarda con la fecha del hito prefijada en
    el nombre (para identificarlo aunque salga de su carpeta); el obra-level
    conserva su nombre. Sin red: session falsa con contenido fijo."""
    from entregables.zip_infoobras import _descargar_a_carpeta

    class _Resp:
        status_code = 200
        headers: dict = {}
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def iter_content(self, chunk_size=1):
            yield b"%PDF-1.4 contenido de prueba"

    class _Sess:
        def get(self, url, params=None, timeout=None, stream=False):
            return _Resp()

    it = {"filename": "doc/val.pdf", "nombre": "documento20250911", "extension": "pdf"}
    ok = _descargar_a_carpeta(_Sess(), it, tmp_path, timeout=5,
                              prefijo_nombre="2025-09 SEPTIEMBRE · ")
    assert ok
    assert [p.name for p in tmp_path.iterdir()] == [
        "2025-09 SEPTIEMBRE · documento20250911.pdf"]

    # obra-level (sin prefijo) conserva su nombre tal cual
    ok2 = _descargar_a_carpeta(
        _Sess(), {"filename": "exp/et.pdf", "nombre": "expediente", "extension": "pdf"},
        tmp_path / "sec", timeout=5)
    assert ok2 and (tmp_path / "sec" / "expediente.pdf").exists()


def test_get_inventario_reintenta_503_pero_no_4xx(monkeypatch):
    """La página de inventario reintenta ante 503 transitorio y devuelve el HTML
    en el intento que responde 200; un 4xx es permanente (no reintenta)."""
    from entregables import zip_infoobras as z

    class _Resp:
        def __init__(self, code, text=""):
            self.status_code, self.text = code, text
        def raise_for_status(self):
            if self.status_code >= 400:
                raise z.requests.HTTPError(f"HTTP {self.status_code}")

    monkeypatch.setattr(z.time, "sleep", lambda *a: None)  # no esperar en el test

    estado = {"n": 0}
    class _Sess503:
        def get(self, url, params=None, timeout=None):
            estado["n"] += 1
            return _Resp(503) if estado["n"] == 1 else _Resp(200, "<html>ok</html>")
    assert z._get_inventario(_Sess503(), 72056, timeout=5, intentos=3) == "<html>ok</html>"
    assert estado["n"] == 2  # cayó la 1ra, respondió la 2da

    estado2 = {"n": 0}
    class _Sess404:
        def get(self, url, params=None, timeout=None):
            estado2["n"] += 1
            return _Resp(404)
    with pytest.raises(z.requests.HTTPError):
        z._get_inventario(_Sess404(), 1, timeout=5, intentos=3)
    assert estado2["n"] == 1  # un 4xx no se reintenta


def test_hoja_profesional_cuadro_emisor_sunat_alt04():
    """La hoja del profesional trae el cuadro EMISOR (SUNAT) al lado, y marca
    ALT04 cuando el certificado se emitió ANTES de la creación de la empresa."""
    import openpyxl
    from entregables.excel_final import construir_hoja_profesional

    ws = openpyxl.Workbook().active
    prof = {"n_prof": 1, "cargo": "ESP", "nombre": "N", "experiencias": [
        {"n": 1, "proyecto": "Obra X", "fecha_inicial": "2019-06-01",
         "fecha_final": "2020-06-30", "fecha_emision": "2019-03-01",  # ANTES de creación
         "ruc_emisor": "20512345678"}]}
    fichas = {(1, 1): {"codigo_infoobras": "111", "cui": "2418877", "valorizaciones": []}}
    sunat = {(1, 1): {"ruc": "20512345678", "razon_social": "CONSORCIO X SAC",
                      "fecha_inscripcion": "2019-06-01", "estado": "ACTIVO"}}
    construir_hoja_profesional(ws, prof, {}, {}, fichas, {}, sunat)

    texto = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "EMISOR DEL CERTIFICADO (SUNAT)" in texto
    assert "CONSORCIO X SAC" in texto
    # creación 2019-06 > emisión 2019-03 → certificado emitido antes de existir la empresa
    assert "ALT04" in texto and "ANTES de" in texto and "creación" in texto


def test_hoja_profesional_emisor_sin_anomalia():
    """Emisor coherente (creado antes del inicio y de la emisión) → ALT04 en verde."""
    import openpyxl
    from entregables.excel_final import construir_hoja_profesional

    ws = openpyxl.Workbook().active
    prof = {"n_prof": 1, "cargo": "ESP", "nombre": "N", "experiencias": [
        {"n": 1, "proyecto": "Obra Y", "fecha_inicial": "2020-01-01",
         "fecha_final": "2021-01-01", "fecha_emision": "2021-02-01",
         "ruc_emisor": "20512345678"}]}
    fichas = {(1, 1): {"cui": "2418877", "valorizaciones": []}}
    sunat = {(1, 1): {"ruc": "20512345678", "razon_social": "EMPRESA OK",
                      "fecha_inscripcion": "2010-05-10", "estado": "ACTIVO",
                      "condicion": "HABIDO", "tipo_contribuyente": "SOCIEDAD ANONIMA CERRADA",
                      "actividades_economicas": ["Principal - 7110 - ARQUITECTURA E INGENIERIA"],
                      "domicilio_fiscal": "AV X 123, LIMA"}}
    construir_hoja_profesional(ws, prof, {}, {}, fichas, {}, sunat)
    texto = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "sin anomalía de antigüedad" in texto
    # objeto social (CIIU) mostrado, con el prefijo "Principal -" limpiado
    assert "Objeto social" in texto and "7110" in texto and "ARQUITECTURA E INGENIERIA" in texto
    assert "Principal -" not in texto
    assert "HABIDO" in texto and "SOCIEDAD ANONIMA CERRADA" in texto


def test_hoja_profesional_representantes_e_historico_del_emisor():
    """Issue #30: el cuadro del emisor lista los representantes legales y responde
    —en el TÍTULO del campo— si estaba habido al emitir y durante la obra; el
    histórico va en un cuadro contiguo (R:U) y el representante de obra se corre
    a W:Z."""
    import openpyxl
    from entregables.excel_final import construir_hoja_profesional

    ws = openpyxl.Workbook().active
    prof = {"n_prof": 1, "cargo": "ESP", "nombre": "N", "experiencias": [
        {"n": 1, "proyecto": "Obra Z", "fecha_inicial": "2021-05-13",
         "fecha_final": "2023-05-06", "fecha_emision": "2023-06-01",
         "ruc_emisor": "20512345678"}]}
    fichas = {(1, 1): {"cui": "2418877", "valorizaciones": [],
                       "representante_obra": {"contratistas": [
                           {"nombre_empresa": "EJECUTORA SAC", "ruc": "20111111111"}]}}}
    sunat = {(1, 1): {
        "ruc": "20512345678", "razon_social": "CONSTRUCTORA X SAC",
        "fecha_inscripcion": "2010-01-05", "estado": "ACTIVO", "condicion": "HABIDO",
        "representantes": [
            {"tipo_documento": "CE", "nro_documento": "001748034",
             "nombre": "ZHANG XIA", "cargo": "APODERADO", "fecha_desde": "2020-08-17"},
            {"tipo_documento": "CE", "nro_documento": "002143422",
             "nombre": "LI WENXUE", "cargo": "APODERADO", "fecha_desde": "2020-12-09"}],
        "historico": {"razones_sociales": [
            {"nombre": "CONSTRUCTORA X EIRL", "fecha_baja": "2016-05-04"}],
            "domicilios": [{"direccion": "JR. AREQUIPA 55", "fecha_baja": "2016-05-04"}]},
        "habido": {
            "emision": {"fecha": "2023-06-01", "condicion": "HABIDO", "ok": True},
            "periodo": {"desde": "2021-05-13", "hasta": "2023-05-06",
                        "tramos": [
                            {"condicion": "HABIDO", "desde": "2021-05-13", "hasta": "2022-06-30"},
                            {"condicion": "NO HABIDO", "desde": "2022-07-01", "hasta": "2022-09-30"}],
                        "tramos_no_habido": [
                            {"condicion": "NO HABIDO", "desde": "2022-07-01", "hasta": "2022-09-30"}],
                        "ok": False}}}}
    construir_hoja_profesional(ws, prof, {}, {}, fichas, {}, sunat)

    celdas = {c.coordinate: str(c.value) for row in ws.iter_rows() for c in row if c.value}
    texto = "\n".join(celdas.values())

    # representantes: todos, con cargo y fecha desde — y SIN cruce con el firmante
    assert "Representantes legales (SUNAT)" in texto
    assert "ZHANG XIA" in texto and "LI WENXUE" in texto
    assert "APODERADO" in texto and "17/08/20" in texto
    assert "ALT12" not in texto and "firmante" not in texto

    # la pregunta va en el TÍTULO del campo, no solo en el valor
    assert any(v.startswith("¿Estaba habido al emitir el certificado?")
               for v in celdas.values())
    assert any(v.startswith("¿Estuvo habido durante la obra?") for v in celdas.values())
    assert "NO — NO HABIDO 01/07/22–30/09/22" in texto

    # cuadro histórico contiguo: arranca en R (col 18), en la misma fila que el emisor
    emisor = next(k for k, v in celdas.items() if v == "EMISOR DEL CERTIFICADO (SUNAT)")
    hist = next(k for k, v in celdas.items() if v == "HISTÓRICO SUNAT DEL EMISOR")
    assert emisor[0] == "M" and hist[0] == "R"
    assert emisor[1:] == hist[1:], "los dos cuadros del emisor arrancan en la misma fila"
    assert "CONSTRUCTORA X EIRL" in texto and "JR. AREQUIPA 55" in texto

    # el representante de obra (InfoObras) se corrió a W:Z
    rep_obra = next(k for k, v in celdas.items() if v == "REPRESENTANTE DE OBRA (InfoObras)")
    assert rep_obra[0] == "W"


def test_hoja_profesional_sin_historico_no_dibuja_el_cuadro():
    """Emisor sin información histórica: no se dibuja el cuadro contiguo (queda
    'no verificable' en el campo del emisor) y no se rompe nada."""
    import openpyxl
    from entregables.excel_final import construir_hoja_profesional

    ws = openpyxl.Workbook().active
    prof = {"n_prof": 1, "cargo": "ESP", "nombre": "N", "experiencias": [
        {"n": 1, "proyecto": "Obra W", "fecha_inicial": "2020-01-01",
         "fecha_final": "2021-01-01", "ruc_emisor": "20512345678"}]}
    sunat = {(1, 1): {"ruc": "20512345678", "razon_social": "EMPRESA OK",
                      "fecha_inscripcion": "2010-05-10", "estado": "ACTIVO",
                      "condicion": "HABIDO", "representantes": [], "historico": None,
                      "habido": None}}
    construir_hoja_profesional(ws, prof, {}, {}, {(1, 1): {"valorizaciones": []}}, {}, sunat)

    texto = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "HISTÓRICO SUNAT DEL EMISOR" not in texto
    assert "Representantes legales" not in texto
    assert "¿Estaba habido al emitir y durante la obra?" in texto
    assert "Sin información histórica en SUNAT" in texto


# ── Visibilidad del bloque de obra (issue #46: nombre de la obra + aviso) ─────


def _texto(ws) -> str:
    return "\n".join(str(c.value) for f in ws.iter_rows() for c in f if c.value)


def _celdas(ws) -> list[str]:
    return [str(c.value) for f in ws.iter_rows() for c in f if c.value is not None]


def _hoja_de(salida, n_prof, cargo):
    return openpyxl.load_workbook(salida)[etiqueta_hoja(n_prof, cargo)]


_ESPEJO_1EXP = {
    "_meta": {}, "postor": {},
    "profesionales": [{
        "n_prof": 1, "cargo": "JEFE", "nombre": "N",
        "experiencias": [{"n": 1, "proyecto": "Construcción del Hospital de EsSalud, Tarapoto",
                          "fecha_inicial": "2019-05-01", "fecha_final": "2019-10-31"}],
    }],
    "resumen_evaluacion": {"factores": []},
}


def test_bloque_de_obra_muestra_el_nombre_primero(tmp_path):
    """Punto 1 del issue #46: sin el nombre, un CUI mal resuelto (obra de OTRA
    región) pasa inadvertido detrás de un código y un monto correctos. El nombre
    es el PRIMER campo del bloque, antes del código."""
    fichas = {(1, 1): {"codigo_infoobras": "104710", "cui": "2405647",
                       "obra_nombre": "INSTALACION DE LOS SERVICIOS DE TOMOGRAFIA, "
                                      "PUERTO MALDONADO, MADRE DE DIOS",
                       "estado": "Finalizado", "valorizaciones": []}}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "n.xlsx", fichas=fichas)
    ws = _hoja_de(salida, 1, "JEFE")

    fila_nom = next(c.row for f in ws.iter_rows() for c in f if c.value == "Nombre de la obra")
    fila_cod = next(c.row for f in ws.iter_rows() for c in f if c.value == "Código InfoObras")
    assert fila_nom < fila_cod                      # primero el nombre
    assert "PUERTO MALDONADO" in _texto(ws)
    # el nombre no cabe en una línea y una celda combinada no auto-ajusta el alto
    assert (ws.row_dimensions[fila_nom].height or 15) > 15


def test_bloque_de_obra_sin_nombre_lo_dice(tmp_path):
    """Ficha sin nombre (InfoObras no lo devolvió): se dice, no se omite la fila —
    la ausencia del dato también es información para el evaluador."""
    fichas = {(1, 1): {"codigo_infoobras": "104710", "cui": "2405647", "valorizaciones": []}}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "sn.xlsx", fichas=fichas)
    assert "no devolvió el nombre" in _texto(_hoja_de(salida, 1, "JEFE"))


def test_desempaquetar_enriquecimiento_lleva_el_nombre_a_la_ficha():
    """El nombre NO viene dentro de `obra_ficha`: viaja suelto (`obra_nombre`) o en
    `obra.nombre_obra`. Sin esta traducción nunca llegaba al Excel."""
    from entregables import desempaquetar_enriquecimiento

    _, _, fichas, _ = desempaquetar_enriquecimiento({
        "1:1": {"obra_nombre": "HOSPITAL X", "obra_ficha": {"cui": "1", "codigo_infoobras": "9"}},
        "1:2": {"obra": {"nombre_obra": "HOSPITAL Y"},          # solo el nombre anidado
                "obra_ficha": {"cui": "2", "codigo_infoobras": "8"}},
    })
    assert fichas[(1, 1)]["obra_nombre"] == "HOSPITAL X"
    assert fichas[(1, 2)]["obra_nombre"] == "HOSPITAL Y"


def test_aviso_de_revision_se_pinta_aunque_haya_obra(tmp_path):
    """Punto 2 del issue #46: antes el aviso era `elif` de la obra, así que si el
    resolver proponía una obra el aviso NUNCA se pintaba (40% de los avisos
    invisibles). Ahora va como banda ENCIMA del bloque de la obra."""
    fichas = {(1, 1): {"codigo_infoobras": "1", "cui": "2405647",
                       "obra_nombre": "OBRA PROPUESTA", "valorizaciones": []}}
    revisiones = {(1, 1): ("los candidatos contradicen el rubro del certificado — confirmar",
                           "elegir candidato o pegar CUI")}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "rev.xlsx",
                                 fichas=fichas, revisiones=revisiones)
    ws = _hoja_de(salida, 1, "JEFE")
    texto = _texto(ws)

    assert "EXPERIENCIA EN REVISIÓN" in texto
    assert "contradicen el rubro" in texto
    assert "Acción: elegir candidato o pegar CUI" in texto
    aviso = next(c for f in ws.iter_rows() for c in f
                 if str(c.value or "").startswith("⚠ EXPERIENCIA EN REVISIÓN"))
    cabecera = next(c for f in ws.iter_rows() for c in f
                    if str(c.value or "").startswith("OBRA EN INFOOBRAS"))
    # misma franja de columnas y ARRIBA del bloque de obra (no como nota al pie)
    assert aviso.column == cabecera.column
    assert aviso.row < cabecera.row
    assert aviso.fill.fgColor.rgb.endswith("FF9999")     # alerta, se ve de lejos


def test_el_aviso_dice_el_motivo_real_y_no_una_narrativa_de_obra(tmp_path):
    """Los ItemRevision los crean 4 etapas y la mayoría no habla de la obra. Un
    aviso que igual grita "la obra no está confirmada, pegue el CUI" sobre un CUI
    correcto es un falso positivo — y un aviso que miente enseña al evaluador a
    ignorar la banda. El motivo y la acción son los del item."""
    fichas = {(1, 1): {"codigo_infoobras": "1", "cui": "2405647",
                       "obra_nombre": "OBRA BIEN IDENTIFICADA", "valorizaciones": []}}
    revisiones = {(1, 1): ("NO CUMPLE pero la(s) experiencia(s) 3 no se contó por una "
                           "fecha sin verificar — confirmarla puede cambiar el veredicto",
                           "verificar la fecha de fin del certificado")}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "rev3.xlsx",
                                 fichas=fichas, revisiones=revisiones)
    texto = _texto(_hoja_de(salida, 1, "JEFE"))

    assert "fecha sin verificar" in texto
    assert "Acción: verificar la fecha de fin del certificado" in texto
    # el motivo se muestra VERBATIM: `.capitalize()` bajaría el veredicto del motor
    assert "NO CUMPLE pero" in texto
    # nada de la narrativa de identidad de obra que antes iba cableada
    assert "PROPUESTA del sistema" not in texto
    assert "pegar el CUI" not in texto


def test_el_aviso_sin_accion_sugerida_no_inventa_una(tmp_path):
    """Un item sin `accion_sugerida` cae a una acción GENÉRICA: mandar a pegar un
    CUI sin saber que el problema es el CUI es fabricar un diagnóstico."""
    fichas = {(1, 1): {"codigo_infoobras": "1", "cui": "2405647",
                       "obra_nombre": "OBRA", "valorizaciones": []}}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "rev4.xlsx", fichas=fichas,
                                 revisiones={(1, 1): "el portal no respondió"})
    texto = _texto(_hoja_de(salida, 1, "JEFE"))
    assert "Acción: revisar esta experiencia en el panel" in texto
    assert "pegar el CUI" not in texto


def test_el_aviso_no_encoge_la_fila_del_titulo_de_la_experiencia(tmp_path):
    """El aviso comparte fila con el título de la experiencia (nombre del proyecto
    VERBATIM, que `banda()` ya dimensionó). Fijar la altura cortaba el nombre justo
    en la experiencia marcada como dudosa: la altura solo puede AUMENTAR."""
    espejo = {
        "_meta": {}, "postor": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "JEFE", "nombre": "N", "experiencias": [{
                "n": 1, "fecha_inicial": "2019-05-01", "fecha_final": "2019-10-31",
                "proyecto": ("MEJORAMIENTO Y AMPLIACIÓN DE LOS SERVICIOS DE SALUD DEL "
                             "ESTABLECIMIENTO DE SALUD I-4 SAN MARTÍN DE PORRES, DISTRITO "
                             "DE MORALES, PROVINCIA DE SAN MARTÍN, DEPARTAMENTO DE SAN "
                             "MARTÍN")}]}],
        "resumen_evaluacion": {"factores": []},
    }
    fichas = {(1, 1): {"codigo_infoobras": "1", "cui": "2405647",
                       "obra_nombre": "OBRA", "valorizaciones": []}}
    sin = generar_excel_final(espejo, tmp_path / "alto_sin.xlsx", fichas=fichas)
    con = generar_excel_final(espejo, tmp_path / "alto_con.xlsx", fichas=fichas,
                              revisiones={(1, 1): ("motivo cualquiera", "hacer algo")})

    def _alto_del_titulo(salida):
        ws = _hoja_de(salida, 1, "JEFE")
        fila = next(c.row for f in ws.iter_rows() for c in f
                    if str(c.value or "").startswith("EXPERIENCIA 1:"))
        return ws.row_dimensions[fila].height or 15

    assert _alto_del_titulo(con) >= _alto_del_titulo(sin) > 15


def test_sin_obra_el_aviso_de_revision_sigue_siendo_el_bloque_completo(tmp_path):
    """Regresión: cuando NO hay obra, el bloque de revisión de siempre (con la
    acción a tomar) sigue ocupando la franja — no se degrada a la banda."""
    revisiones = {(1, 1): "el certificado no cita CUI"}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "rev2.xlsx", revisiones=revisiones)
    texto = _texto(_hoja_de(salida, 1, "JEFE"))
    assert "EXPERIENCIA EN REVISIÓN" in texto
    # aquí sí no hubo cruce, así que la acción por defecto es de identidad de obra
    assert "pegar el CUI correcto en el panel" in texto
    assert "OBRA EN INFOOBRAS" not in texto


def test_sin_obra_el_bloque_usa_la_accion_del_item_si_la_trae(tmp_path):
    """…pero si el item trae su propia acción, manda la del item."""
    revisiones = {(1, 1): ("no se pudo verificar en InfoObras — veredicto provisional",
                           "reintentar la consulta o verificar la obra a mano")}
    salida = generar_excel_final(_ESPEJO_1EXP, tmp_path / "rev5.xlsx", revisiones=revisiones)
    texto = _texto(_hoja_de(salida, 1, "JEFE"))
    assert "Acción: reintentar la consulta o verificar la obra a mano" in texto
    assert "pegar el CUI" not in texto


def test_una_obra_repetida_se_detalla_completa_en_cada_experiencia(tmp_path):
    """El bloque de UNA EXPERIENCIA se pinta siempre completo, aunque la obra ya
    esté detallada arriba: es la evidencia de esa experiencia, en la franja pegada
    a su cuadro de hitos. El colapso (issue #46 punto 6) solo aplica a las
    SubExperiencias repetidas de un cert multi-obra.

    Además cubre el caso que hizo RETIRAR la primera versión del colapso: mismo
    CUI, obras DISTINTAS (la ventana del certificado desambigua). La clave es
    (CUI, código InfoObras), así que ni siquiera son la misma obra."""
    espejo = {
        "_meta": {}, "postor": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "JEFE", "nombre": "N", "experiencias": [
                {"n": 1, "proyecto": "Tramo 1", "fecha_inicial": "2019-01-01",
                 "fecha_final": "2019-06-30"},
                {"n": 2, "proyecto": "Tramo 2", "fecha_inicial": "2021-01-01",
                 "fecha_final": "2021-12-31"}]}],
        "resumen_evaluacion": {"factores": []},
    }
    fichas = {
        (1, 1): {"cui": "2418877", "codigo_infoobras": "111", "obra_nombre": "OBRA A",
                 "valorizaciones": [{"anio": 2019, "mes": 3, "estado": "En ejecución"}]},
        # mismo CUI, OTRA obra y otro periodo: sin valorizaciones registradas
        (1, 2): {"cui": "2418877", "codigo_infoobras": "999", "obra_nombre": "OBRA B",
                 "valorizaciones": []},
    }
    salida = generar_excel_final(espejo, tmp_path / "dup.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    texto = "\n".join(celdas)

    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert "999" in texto and "OBRA B" in texto          # la 2ª ficha no se borró
    assert "YA DETALLADA" not in texto


def test_la_misma_obra_en_dos_experiencias_no_colapsa_nunca(tmp_path):
    """Ni siquiera con ficha IDÉNTICA: dos experiencias distintas (dos
    certificados) se pintan completas cada una. Colapsar la franja F:K de una
    experiencia dejaría su cuadro de hitos sin la evidencia al lado."""
    espejo = {
        "_meta": {}, "postor": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "JEFE", "nombre": "N", "experiencias": [
                {"n": 1, "proyecto": "Orden 1", "fecha_inicial": "2019-01-01",
                 "fecha_final": "2019-06-30"},
                {"n": 2, "proyecto": "Orden 2", "fecha_inicial": "2019-01-01",
                 "fecha_final": "2019-06-30"}]}],
        "resumen_evaluacion": {"factores": []},
    }
    ficha = {"cui": "2418877", "codigo_infoobras": "111", "obra_nombre": "OBRA A",
             "estado": "Finalizado", "valorizaciones": []}
    salida = generar_excel_final(espejo, tmp_path / "dup2.xlsx",
                                 fichas={(1, 1): dict(ficha), (1, 2): dict(ficha)})
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)


# ── Colapso de sub-obras repetidas (issue #46 punto 6, rediseñado) ────────────
# Fixtures MODELADAS COMO EL PIPELINE REAL: `_enriquecer_subobra` SIEMPRE mete
# las valorizaciones de la obra dentro de `ficha`, así que una sub-obra con
# `valorizaciones: []` solo representa a una obra que InfoObras no valorizó.

def _espejo_multi(n_exps: int, periodos: list[tuple[str, str]]) -> dict:
    return {
        "_meta": {}, "postor": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "JEFE", "nombre": "N",
            "experiencias": [
                {"n": i, "proyecto": f"Orden de servicio {i}",
                 "fecha_inicial": periodos[i - 1][0], "fecha_final": periodos[i - 1][1]}
                for i in range(1, n_exps + 1)],
        }],
        "resumen_evaluacion": {"factores": []},
    }


def _sub(cui, cod, nombre, vals, obra_nombre=None, **extra):
    """Sub-obra como la deja `_enriquecer_subobra`. `nombre` es el proyecto tal
    como lo DECLARA el certificado y `obra_nombre` el de InfoObras: en los certs
    multi-obra reales pueden diferir (dos establecimientos declarados aparte que
    son una sola obra registrada)."""
    return {"estado": "resuelto", "cui": cui, "via": "CODIGO", "proyecto": nombre,
            "fecha_inicial": None, "fecha_final": None,
            "ficha": {"cui": cui, "codigo_infoobras": cod,
                      "obra_nombre": obra_nombre or nombre,
                      "estado": "En Ejecución", "monto": 100.0,
                      "fecha_inicio": "2018-01-01", "fecha_fin": "2019-01-01",
                      "modificaciones_plazo": [], "valorizaciones": vals},
            **extra}


def test_subobras_identicas_de_ordenes_consecutivas_se_colapsan(tmp_path):
    """El caso real del job 95af90f1578e / P3. Arquit: 4 órdenes de servicio
    consecutivas del mismo contrato, cada una certificando las MISMAS 3 obras →
    12 bloques de detalle idénticos. Se pintan 3 y las 9 repeticiones quedan como
    puntero a la fila donde está el detalle."""
    periodos = [("2025-06-11", "2025-08-31"), ("2025-09-02", "2025-12-30"),
                ("2026-01-28", "2026-03-31"), ("2026-04-07", "2026-06-30")]
    espejo = _espejo_multi(4, periodos)
    subs = [_sub("2377760", "92111", "OBRA PONGO", []),
            _sub("2377762", "92112", "OBRA CHAZUTA", []),
            _sub("2377764", "92168", "OBRA TABALOSOS", [])]
    fichas = {(1, i): {"sub_obras": [dict(s) for s in subs]} for i in range(1, 5)}
    salida = generar_excel_final(espejo, tmp_path / "multi.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))

    # las 12 SubExperiencias SIGUEN estando: cada certificado declara esas obras y
    # ese hecho es evidencia de ESE certificado — lo que se colapsa es el detalle.
    assert sum(1 for v in celdas if v.startswith("SubExperiencia ")) == 12
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 3
    assert sum(1 for v in celdas if v.startswith("OBRA YA DETALLADA")) == 9
    texto = "\n".join(celdas)
    # el puntero nombra la obra, dónde está el detalle y QUÉ hay allá
    assert "código InfoObras 92111" in texto
    assert "SubExperiencia 1.1" in texto
    assert "InfoObras no registra valorizaciones de esta obra" in texto


def test_dos_subobras_del_mismo_cert_al_mismo_cui_colapsan_una_vez(tmp_path):
    """Caso real P1. Jefe Superv exp 3: dos sub-obras declaradas con nombres
    distintos (C.S. TONGOD y P.S. PISIT) resuelven a la MISMA obra de InfoObras
    → un detalle y un puntero. El nombre declarado de la 2ª NO se pierde."""
    espejo = _espejo_multi(1, [("2024-06-12", "2024-09-16")])
    vals = [{"anio": 2024, "mes": 10, "estado": "En ejecución",
             "fisico_real": 71.6, "valorizado_real": 3142323.82, "docs": 1}]
    obra = "MEJORAMIENTO ... C.S TONGOD Y P.S PISIT ... SAN MIGUEL - CAJAMARCA"
    fichas = {(1, 1): {"sub_obras": [
        _sub("2440017", "504249", "C.S. TONGOD", list(vals), obra_nombre=obra),
        _sub("2440017", "504249", "P.S. PISIT", list(vals), obra_nombre=obra)]}}
    salida = generar_excel_final(espejo, tmp_path / "tongod.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    texto = "\n".join(celdas)

    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 1
    assert sum(1 for v in celdas if v.startswith("OBRA YA DETALLADA")) == 1
    # los DOS nombres declarados siguen visibles (las cabeceras no se colapsan):
    # que dos establecimientos declarados sean una sola obra es justo lo que el
    # evaluador tiene que ver.
    assert "P.S. PISIT" in texto and "C.S. TONGOD" in texto
    assert "su única valorización" in texto


def test_no_colapsa_si_el_periodo_resaltado_cambia(tmp_path):
    """Dos apariciones de la MISMA obra con valorizaciones pero periodos de
    certificado distintos pintan resaltados distintos: son información distinta →
    ambas completas. Es la condición de igualdad de periodo."""
    espejo = _espejo_multi(2, [("2019-03-01", "2019-03-31"),
                               ("2019-05-01", "2019-05-31")])
    vals = [{"anio": 2019, "mes": 3, "estado": "En ejecución"},
            {"anio": 2019, "mes": 5, "estado": "En ejecución"}]
    fichas = {(1, i): {"sub_obras": [_sub("2418877", "111", "OBRA A", list(vals))]}
              for i in (1, 2)}
    salida = generar_excel_final(espejo, tmp_path / "per.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)


def test_no_colapsa_por_periodo_distinto_aunque_no_cambie_el_resaltado(tmp_path):
    """La igualdad de periodo se exige en crudo, no solo por sus efectos: aquí
    ninguna valorización cae dentro de ninguna de las dos ventanas, así que el
    resaltado sería el mismo — y aun así los periodos distintos bloquean el
    colapso."""
    espejo = _espejo_multi(2, [("2021-01-01", "2021-01-31"),
                               ("2022-06-01", "2022-06-30")])
    vals = [{"anio": 2019, "mes": 3, "estado": "En ejecución"}]
    fichas = {(1, i): {"sub_obras": [_sub("2418877", "111", "OBRA A", list(vals))]}
              for i in (1, 2)}
    salida = generar_excel_final(espejo, tmp_path / "perb.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)


def test_no_colapsa_si_cambian_los_montos_de_las_valorizaciones(tmp_path):
    """Mismo número de valorizaciones no es el mismo contenido. El portal de
    InfoObras es inestable (tablas que varían entre corridas): si dos fichas de
    la misma obra traen importes distintos, esa diferencia es justo lo que hay
    que ver, no lo que hay que esconder tras un puntero."""
    espejo = _espejo_multi(2, [("2019-01-01", "2019-12-31"),
                               ("2019-01-01", "2019-12-31")])
    fichas = {
        (1, 1): {"sub_obras": [_sub("2418877", "111", "OBRA A", [
            {"anio": 2019, "mes": 3, "estado": "En ejecución", "valorizado_real": 100.0}])]},
        (1, 2): {"sub_obras": [_sub("2418877", "111", "OBRA A", [
            {"anio": 2019, "mes": 3, "estado": "En ejecución", "valorizado_real": 999.0}])]},
    }
    salida = generar_excel_final(espejo, tmp_path / "mon.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)
    assert 999.0 in [c.value for f in _hoja_de(salida, 1, "JEFE").iter_rows() for c in f]


@pytest.mark.parametrize("campo,valor", [
    ("obra_nombre", "OTRA OBRA"),
    ("estado", "Finalizado"),
    ("monto", 999.0),
    ("fecha_inicio", "2010-01-01"),
    ("fecha_fin", "2030-12-31"),
    ("modificaciones_plazo", [{"tipo": "Ampliación", "dias": 30}]),
    ("representante_obra", {"contratistas": [{"ruc": "20100000001",
                                              "nombre_empresa": "CONSTRUCTORA X"}]}),
])
def test_no_colapsa_si_difiere_cualquier_campo_de_la_ficha(tmp_path, campo, valor):
    """Cada campo que el bloque pinta entra en la huella. Si la 2ª aparición de la
    misma obra trae algo distinto — otro estado, otro monto, otras fechas de obra,
    una ampliación de plazo, el representante — se pinta completa: esa diferencia
    es información, y esconderla tras "igual que arriba" sería mentir."""
    espejo = _espejo_multi(2, [("2019-01-01", "2019-12-31"),
                               ("2019-01-01", "2019-12-31")])
    sub2 = _sub("2418877", "111", "OBRA A", [])
    sub2["ficha"][campo] = valor
    fichas = {(1, 1): {"sub_obras": [_sub("2418877", "111", "OBRA A", [])]},
              (1, 2): {"sub_obras": [sub2]}}
    salida = generar_excel_final(espejo, tmp_path / f"{campo}.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)


def test_no_colapsa_si_una_ficha_es_expediente_y_la_otra_obra(tmp_path):
    """Misma obra, mismo periodo, pero una aparición trae la verificación del
    EXPEDIENTE TÉCNICO (MEF) y la otra no: son bloques que dicen cosas distintas."""
    espejo = _espejo_multi(2, [("2019-01-01", "2019-12-31"),
                               ("2019-01-01", "2019-12-31")])
    verif = {"contratista": {"valor": "CONSORCIO X", "veredicto": "ok"},
             "contrato": {"numero": "007"}, "resolucion": {"numero": "R-1"},
             "cui_confirmado": "2418877", "fuentes": ["MEF"]}
    sub_con = _sub("2418877", "111", "OBRA A", [])
    sub_con["ficha"]["verificacion_expediente"] = verif
    fichas = {(1, 1): {"sub_obras": [sub_con]},
              (1, 2): {"sub_obras": [_sub("2418877", "111", "OBRA A", [])]}}
    salida = generar_excel_final(espejo, tmp_path / "exp.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    texto = "\n".join(celdas)
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)
    assert "VERIFICACIÓN SEACE / MEF" in texto and "CONSORCIO X" in texto


def test_colapso_de_bloque_de_expediente_dice_que_es_un_expediente(tmp_path):
    """Cuando lo repetido SÍ es un bloque de expediente, el puntero lo nombra:
    nunca dice "valorizaciones" donde no las hay."""
    espejo = _espejo_multi(1, [("2019-01-01", "2019-12-31")])
    verif = {"contratista": {"valor": "CONSORCIO X", "veredicto": "ok"},
             "contrato": {"numero": "007"}, "resolucion": {"numero": "R-1"},
             "cui_confirmado": "2418877", "fuentes": ["MEF"]}

    def _con_verif(nombre):
        s = _sub("2418877", "111", nombre, [], obra_nombre="OBRA A")
        s["ficha"]["verificacion_expediente"] = dict(verif)
        return s

    fichas = {(1, 1): {"sub_obras": [_con_verif("Tramo A"), _con_verif("Tramo B")]}}
    salida = generar_excel_final(espejo, tmp_path / "expc.xlsx", fichas=fichas)
    texto = _texto(_hoja_de(salida, 1, "JEFE"))
    assert "OBRA YA DETALLADA" in texto
    assert "la verificación del expediente técnico" in texto
    assert "valorizaciones de esta obra" not in texto


def test_no_colapsa_si_una_ficha_trae_valorizaciones_y_la_otra_no(tmp_path):
    """La trampa que hundió la primera versión: una ficha SIN valorizaciones
    referida a otra que SÍ las tiene renderiza una AUSENCIA de dato como puntero
    a evidencia. Aunque sea la misma obra y el mismo periodo, no se colapsa."""
    espejo = _espejo_multi(2, [("2019-01-01", "2019-12-31"),
                               ("2019-01-01", "2019-12-31")])
    fichas = {
        (1, 1): {"sub_obras": [_sub("2418877", "111", "OBRA A",
                                    [{"anio": 2019, "mes": 3, "estado": "En ejecución"}])]},
        (1, 2): {"sub_obras": [_sub("2418877", "111", "OBRA A", [])]},
    }
    salida = generar_excel_final(espejo, tmp_path / "vac.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)
    assert "Sin valorizaciones registradas en InfoObras" in "\n".join(celdas)


def test_no_colapsa_dos_obras_distintas_del_mismo_cui(tmp_path):
    """Mismo CUI, códigos InfoObras distintos = obras distintas (`_fetch_obra`
    desambigua por la ventana del certificado). La clave lleva los dos."""
    espejo = _espejo_multi(2, [("2019-01-01", "2019-12-31"),
                               ("2019-01-01", "2019-12-31")])
    fichas = {
        (1, 1): {"sub_obras": [_sub("2418877", "111", "OBRA A", [])]},
        (1, 2): {"sub_obras": [_sub("2418877", "999", "OBRA B", [])]},
    }
    salida = generar_excel_final(espejo, tmp_path / "cui.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert "OBRA B" in "\n".join(celdas)
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)


def test_sin_codigo_infoobras_no_hay_identidad_y_no_se_colapsa(tmp_path):
    """Sin código de obra no hay identidad comprobable: ausencia de dato NUNCA
    habilita el colapso."""
    espejo = _espejo_multi(2, [("2019-01-01", "2019-12-31"),
                               ("2019-01-01", "2019-12-31")])
    sub = _sub("2418877", None, "OBRA A", [])
    fichas = {(1, i): {"sub_obras": [dict(sub)]} for i in (1, 2)}
    salida = generar_excel_final(espejo, tmp_path / "sincod.xlsx", fichas=fichas)
    celdas = _celdas(_hoja_de(salida, 1, "JEFE"))
    assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 2
    assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas)


def test_el_registro_de_obras_pintadas_es_por_hoja_de_profesional(tmp_path):
    """Un "ver SubExperiencia 1.1, fila 92" que apunte a la pestaña de OTRO
    profesional es inservible: cada hoja arranca de cero."""
    espejo = {
        "_meta": {}, "postor": {},
        "profesionales": [
            {"n_prof": 1, "cargo": "JEFE", "nombre": "Uno",
             "experiencias": [{"n": 1, "proyecto": "Orden 1",
                               "fecha_inicial": "2019-01-01", "fecha_final": "2019-12-31"}]},
            {"n_prof": 2, "cargo": "ESTRUCTURAS", "nombre": "Dos",
             "experiencias": [{"n": 1, "proyecto": "Orden 1",
                               "fecha_inicial": "2019-01-01", "fecha_final": "2019-12-31"}]},
        ],
        "resumen_evaluacion": {"factores": []},
    }
    sub = _sub("2418877", "111", "OBRA A", [])
    fichas = {(1, 1): {"sub_obras": [dict(sub)]}, (2, 1): {"sub_obras": [dict(sub)]}}
    salida = generar_excel_final(espejo, tmp_path / "hojas.xlsx", fichas=fichas)
    for n_prof, cargo in ((1, "JEFE"), (2, "ESTRUCTURAS")):
        celdas = _celdas(_hoja_de(salida, n_prof, cargo))
        assert sum(1 for v in celdas if v.startswith("OBRA EN INFOOBRAS")) == 1, cargo
        assert not any(v.startswith("OBRA YA DETALLADA") for v in celdas), cargo


def test_el_colapso_no_pisa_la_cobertura_de_la_subexperiencia(tmp_path):
    """La cobertura del tiempo es POR sub-experiencia (la ventana del certificado
    contra las valorizaciones), no de la obra: se sigue pintando aunque el detalle
    de la obra se haya colapsado."""
    espejo = _espejo_multi(1, [("2019-01-01", "2019-12-31")])
    fichas = {(1, 1): {"sub_obras": [
        _sub("2418877", "111", "OBRA A", []),
        _sub("2418877", "111", "OBRA A", [],
             cobertura={"cubierto": False, "pct": 40, "cert_ini": "2019-01-01",
                        "cert_fin": "2019-12-31"})]}}
    salida = generar_excel_final(espejo, tmp_path / "cob.xlsx", fichas=fichas)
    texto = _texto(_hoja_de(salida, 1, "JEFE"))
    assert "OBRA YA DETALLADA" in texto
    assert "cobertura 40%" in texto and "PARCIAL, revisar" in texto


# ── Hoja CLAUDE: formato de Manuel (blanco + verde/rojo semántico) ─────────────

def test_clasificar_veredicto_polaridad():
    """Verde/rojo según polaridad: en campos donde "SÍ" es malo (¿anterior a
    colegiatura?, ¿antes de culminar?) el SÍ debe salir ROJO, no verde."""
    from scripts.generar_excel import _fill_veredicto

    def color(v, pol):
        f = _fill_veredicto(v, pol)
        return None if f is None else f.fgColor.rgb[-6:]

    assert color("SÍ", "pos") == "C6EFCE"          # verde
    assert color("NO", "pos") == "FFC7CE"          # rojo
    assert color("SÍ", "neg") == "FFC7CE"          # SÍ malo → rojo
    assert color("NO", "neg") == "C6EFCE"          # NO bueno → verde
    assert color("CUMPLE", "pos") == "C6EFCE"
    assert color("NO CUMPLE", "pos") == "FFC7CE"
    assert color("POR VERIFICAR (x)", "pos") == "FFF2CC"   # amarillo
    assert color("NO APLICA", "pos") == "FFF2CC"
    assert color("texto libre", "pos") is None     # sin color
    assert color("", "pos") is None


def test_hoja_claude_blanca_con_veredictos(tmp_path):
    """La hoja CLAUDE es blanca dominante con verde/rojo SOLO en veredictos."""
    espejo = {
        "_meta": {"analisis_id": "t", "concurso": "C", "postor": "P"},
        "postor": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "ESP", "nombre": "N", "profesion_valida": "SÍ",
            "cumple": "CUMPLE — 5 años efectivos",
            "experiencias": [{
                "n": 1, "proyecto": "Obra", "fecha_inicial": "2020-01-01",
                "fecha_final": "2021-01-01", "cargo_valido_emitir": "SÍ",
                "anterior_colegiatura": "SÍ", "cert_antes_culminar": "NO",
                "dias": 367, "folio": "1",
            }],
        }],
        "resumen_evaluacion": {"factores": [], "puntaje_total": None},
    }
    salida = generar_excel_final(espejo, tmp_path / "f.xlsx", {})
    ws = openpyxl.load_workbook(salida)["CLAUDE"]
    from collections import Counter
    fills = Counter(c.fill.fgColor.rgb[-6:] if (c.fill and c.fill.patternType) else None
                    for row in ws.iter_rows() for c in row)
    # blanco dominante; verde y rojo presentes (veredictos)
    assert fills[None] > sum(v for k, v in fills.items() if k)
    assert fills["C6EFCE"] >= 2     # ¿profesión? + ¿válido emitir? + ¿cumple? → verde
    assert fills["FFC7CE"] >= 1     # ¿anterior a colegiatura? SÍ → rojo
    # ya NO se usa el amarillo de "llenado por Claude" en cada celda
    assert "Verde = cumple" in "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value)


def test_excel_final_requisitos_lista_no_crashea(tmp_path):
    """Regresión: `requisitos.*` puede llegar como lista (el espejo permite
    record(any)); openpyxl no escribe listas en celdas. El generador coerce
    list/dict→string en vez de reventar (bloque del TDR por profesional, mejora A)."""
    espejo = {
        "_meta": {},
        "profesionales": [{
            "n_prof": 1, "cargo": "Esp", "nombre": "JUAN", "colegiatura": "CIP1",
            "requisitos": {
                "cargos_validos": ["Residente de obra", "Jefe de obra", "Supervisor de obra"],
                "tipo_experiencia_valida": "5 años",
                "tipo_obra_valida": ["Hospital II-1", "Hospital II-2"],  # LISTA → antes crasheaba
            },
            "experiencias": [{"n": 1, "proyecto": "H", "fecha_inicial": "2019-01-01",
                              "fecha_final": "2020-01-01", "entidad_emisora": "GR",
                              "cargo_ocupado": "Sup"}],
        }],
        "resumen_evaluacion": {"factores": []}, "postor": {},
    }
    salida = generar_excel_final(espejo, tmp_path / "lista.xlsx")   # NO debe crashear
    ws = next(s for s in openpyxl.load_workbook(salida).worksheets if s.title.startswith("P1"))
    rows = {str(row[0].value or ""): row[1].value for row in ws.iter_rows(min_col=1, max_col=2)}
    cargos = next((rows[k] for k in rows if "CARGOS VÁLIDOS" in k), None)
    assert cargos == "Residente de obra; Jefe de obra; Supervisor de obra"


# ── Factor de evaluación embebido al final de la hoja (issue #32) ────────────

_ESPEJO_FACTOR = {
    "_meta": {"pagina_factores": 48},
    "postor": {},
    "profesionales": [{
        "n_prof": 1, "cargo": "ESPECIALISTA EN ARQUITECTURA", "nombre": "JUANA",
        "colegiatura": "CAP 1",
        "experiencias": [{"n": 1, "proyecto": "Hospital X", "fecha_inicial": "2019-01-01",
                          "fecha_final": "2020-01-01", "entidad_emisora": "GR",
                          "cargo_ocupado": "Arquitecta"}],
    }],
    "resumen_evaluacion": {"factores": []},
}


def _pdf_una_pagina(destino: Path, texto: str) -> Path:
    """PDF mínimo de 1 página, con PyMuPDF — el mismo motor con el que el generador
    rasteriza los recortes."""
    fitz = pytest.importorskip("fitz")
    doc = fitz.open()
    doc.new_page().insert_text((72, 100), texto)
    doc.save(str(destino))
    doc.close()
    return destino


def _fila_de(ws, texto: str) -> int:
    """Nº de fila cuya columna A empieza con `texto` (0 si no está)."""
    return next((c.row for (c,) in ws.iter_rows(min_col=1, max_col=1)
                 if str(c.value or "").startswith(texto)), 0)


def test_excel_final_embebe_factor_al_final(tmp_path):
    """El recorte del Cuadro de Factores cierra la hoja del profesional, DESPUÉS de
    la PARTE 5 — y sin pisarla."""
    factor = _pdf_una_pagina(tmp_path / "factor.pdf", "4.2 FACTORES DE EVALUACION")
    salida = generar_excel_final(_ESPEJO_FACTOR, tmp_path / "factor.xlsx",
                                 certificados={(1, "FACTOR"): factor})
    ws = next(s for s in openpyxl.load_workbook(salida).worksheets if s.title.startswith("P1"))

    fila_banda = _fila_de(ws, "FACTOR DE EVALUACIÓN")
    fila_anios = _fila_de(ws, "AÑOS EFECTIVOS")
    # (1) el recorte está, y (2) va DESPUÉS de la PARTE 5 (al final, no al inicio)
    assert fila_banda > 0 and fila_anios > 0
    assert fila_banda > fila_anios
    # (3) REGRESIÓN del `r += 2`: la PARTE 5 deja `r` EN su última fila, así que sin
    # el salto la banda se mergea encima y openpyxl la absorbe SIN error — el Excel
    # sale válido pero con "AÑOS EFECTIVOS" pisado y el número huérfano en la col. S.
    # Por eso NO basta con contar imágenes: hay que exigir que la fila sobreviva.
    assert isinstance(ws.cell(fila_anios, 19).value, (int, float))
    assert len(ws._images) == 1


def test_excel_final_sin_factor_no_cambia_la_hoja(tmp_path):
    """Corrida sin el recorte (agent-bases no capturó la página): la hoja sale igual
    que siempre, sin imagen ni banda — no se rompe la generación."""
    salida = generar_excel_final(_ESPEJO_FACTOR, tmp_path / "sin_factor.xlsx")
    ws = next(s for s in openpyxl.load_workbook(salida).worksheets if s.title.startswith("P1"))
    assert _fila_de(ws, "FACTOR DE EVALUACIÓN") == 0
    assert _fila_de(ws, "AÑOS EFECTIVOS") > 0
    assert not ws._images


def test_regenerar_excel_final_conserva_los_embeds(tmp_path):
    """Regresión: `regenerar_excel_final` no pasaba `certificados`, así que el
    backfill dejaba el entregable SIN ninguna imagen (constancias, TDR, Anexo). Ahora
    mapea la carpeta `certs/` hermana del Excel."""
    from entregables import regenerar_excel_final

    job = tmp_path / "job-1"
    (job / "certs").mkdir(parents=True)
    _pdf_una_pagina(job / "certs" / "P1_FACTOR.pdf", "4.2 FACTORES")
    _pdf_una_pagina(job / "certs" / "P1_E1.pdf", "CONSTANCIA")

    salida = regenerar_excel_final(_ESPEJO_FACTOR, {}, job / "final.xlsx")
    ws = next(s for s in openpyxl.load_workbook(salida).worksheets if s.title.startswith("P1"))
    assert len(ws._images) == 2        # constancia de la experiencia + factor
    fila_anios = _fila_de(ws, "AÑOS EFECTIVOS")
    assert fila_anios > 0
    assert _fila_de(ws, "FACTOR DE EVALUACIÓN") > fila_anios


def test_generar_excel_parte_1_documento_o_descripcion(tmp_path):
    espejo = {
        "_meta": {"analisis_id": "test", "concurso": "C", "postor": "P"},
        "postor": {
            "formularios": [
                {"anexo": "Anexo 1", "documento": "Documento Vigente", "folio": "10"},
                {"anexo": "Anexo 2", "descripcion": "Descripcion Trujillo", "folio": "20"}
            ]
        },
        "profesionales": [],
        "resumen_evaluacion": {"factores": []}
    }
    salida = tmp_path / "claude.xlsx"
    generar_excel(espejo, salida)
    wb = openpyxl.load_workbook(salida)
    ws = wb["CLAUDE"]
    texto = "\n".join(str(c.value) for fila in ws.iter_rows() for c in fila if c.value)
    assert "Documento Vigente" in texto
    assert "Descripcion Trujillo" in texto

