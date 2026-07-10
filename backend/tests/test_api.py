"""
Tests de la API HTTP (FastAPI + TestClient) — el flujo completo que la demo
recorre: crear concurso → subir análisis → pipeline corre → ver job/extracción/
resumen → resolver revisión → descargar Excel final y ZIP. Sin red externa
(etapas del esqueleto: ingesta real + stubs).
"""
from __future__ import annotations

import importlib
import io
import json
import zipfile

import pytest


ESPEJO = {
    "_meta": {"analisis_id": "api-demo-001", "concurso": "CP-API/2026",
              "postor": "POSTOR API DEMO"},
    "postor": {},
    "profesionales": [
        {"n_prof": 1, "cargo": "JEFE DE SUPERVISIÓN", "nombre": "Profesional Uno",
         "colegiatura": "CIP 12345", "cumple": "SÍ — 4.51 años",
         "total": {"dias": 1645, "anios": 4.51},
         "experiencias": [
             {"n": 1, "proyecto": "Obra API A", "entidad_emisora": "Entidad A",
              "fecha_inicial": "2021-05-13", "fecha_final": "2023-05-06",
              "dias": 724, "cui": "2418877", "folio": "100"},
         ]},
    ],
    "resumen_evaluacion": {"factores": [{"factor": "A", "puntaje": 55}],
                           "puntaje_total": 55},
}


@pytest.fixture()
def cliente(tmp_path, monkeypatch):
    """App fresca con datos en tmp_path (aisla cada test)."""
    monkeypatch.setenv("PIVOTE_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("PIVOTE_ETAPAS", "esqueleto")  # sin red en tests
    import api.app as modulo
    importlib.reload(modulo)
    from fastapi.testclient import TestClient
    return TestClient(modulo.app)


def _subir(cliente, concurso_id: str):
    return cliente.post("/api/pivote/analizar", data={"concurso_id": concurso_id},
                        files={
                            "espejo": ("espejo.json", io.BytesIO(
                                json.dumps(ESPEJO).encode("utf-8")), "application/json"),
                            "excel": ("claude.xlsx", io.BytesIO(b"PK demo"), "application/octet-stream"),
                        })


def test_flujo_completo_de_la_demo(cliente):
    # 1 · crear concurso
    r = cliente.post("/api/pivote/concursos",
                     json={"nomenclatura": "CP-API/2026", "entidad": "Entidad Demo"})
    assert r.status_code == 201
    cid = r.json()["concurso_id"]

    # 2 · subir análisis (el pipeline corre como tarea de fondo del TestClient)
    r = _subir(cliente, cid)
    assert r.status_code == 201
    job_id = r.json()["job_id"]

    # 3 · job procesado: 8 pasos OK (esqueleto), estado completado
    r = cliente.get(f"/api/pivote/jobs/{job_id}")
    assert r.status_code == 200
    job = r.json()
    assert job["estado"] == "completado"
    assert len(job["etapas"]) == 8
    assert job["origen"] == "dropzone"

    # 4 · el concurso lista el job
    r = cliente.get(f"/api/pivote/concursos/{cid}")
    assert r.json()["jobs"][0]["job_id"] == job_id
    r = cliente.get("/api/pivote/concursos")
    assert r.json()[0]["n_jobs"] == 1

    # 5 · extracción (profesionales/experiencias para el panel)
    r = cliente.get(f"/api/pivote/jobs/{job_id}/espejo")
    profs = r.json()["profesionales"]
    assert profs[0]["cargo"] == "JEFE DE SUPERVISIÓN"
    assert profs[0]["experiencias"][0]["cui"] == "2418877"

    # 6 · resumen (veredictos + factores)
    r = cliente.get(f"/api/pivote/jobs/{job_id}/resumen")
    res = r.json()
    assert res["puntaje_total"] == 55
    assert res["veredictos"][0]["cumple_claude"].startswith("SÍ")

    # 7 · Excel final REAL descargable (CLAUDE + Base de Datos + hoja del prof)
    r = cliente.get(f"/api/pivote/jobs/{job_id}/excel")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:2] == ["CLAUDE", "Base de Datos"]
    assert any(n.startswith("P1 ") for n in wb.sheetnames)

    # 8 · ZIP InfoObras descargable (árbol de 4 niveles, sin docs aún)
    r = cliente.get(f"/api/pivote/jobs/{job_id}/zip")
    assert r.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    assert any(n.endswith("SIN_DOCUMENTOS.txt") for n in zf.namelist())


def test_espejo_invalido_rechazado_con_errores_claros(cliente):
    cid = cliente.post("/api/pivote/concursos", json={"nomenclatura": "X"}).json()["concurso_id"]
    roto = {"_meta": {"analisis_id": "x"}, "postor": {}, "profesionales": []}
    r = cliente.post("/api/pivote/analizar", data={"concurso_id": cid},
                     files={
                         "espejo": ("e.json", io.BytesIO(json.dumps(roto).encode()), "application/json"),
                         "excel": ("c.xlsx", io.BytesIO(b"PK"), "application/octet-stream"),
                     })
    assert r.status_code == 422
    detalle = r.json()["detail"]
    assert detalle["error"].startswith("el archivo de datos")
    assert detalle["errores"]  # lista campo a campo


def test_decision_de_alerta_persiste(cliente):
    cid = cliente.post("/api/pivote/concursos", json={"nomenclatura": "Y"}).json()["concurso_id"]
    job_id = _subir(cliente, cid).json()["job_id"]
    # sin alertas en el esqueleto: decidir una inexistente igual persiste el registro
    r = cliente.post(f"/api/pivote/jobs/{job_id}/alertas",
                     json={"alerta_id": "al-0", "relevante": False, "razon": "verificado a mano"})
    assert r.status_code == 200


def test_revision_inexistente_404(cliente):
    cid = cliente.post("/api/pivote/concursos", json={"nomenclatura": "Z"}).json()["concurso_id"]
    job_id = _subir(cliente, cid).json()["job_id"]
    r = cliente.post(f"/api/pivote/jobs/{job_id}/revision",
                     json={"n_prof": 9, "n_exp": 9, "cui": "123"})
    assert r.status_code == 404


# ── Descarga por un solo CUI (endpoint suelto) ───────────────────────────────

def test_descargar_cui_valida_entrada(cliente):
    assert cliente.post("/api/pivote/descargar-cui", json={}).status_code == 400
    assert cliente.post("/api/pivote/descargar-cui", json={"cui": "abc"}).status_code == 400


def test_descargar_cui_estado_y_zip_inexistentes_404(cliente):
    assert cliente.get("/api/pivote/descargar-cui/nope").status_code == 404
    assert cliente.get("/api/pivote/descargar-cui/nope/zip").status_code == 404


def test_descargar_cui_flujo_feliz(cliente, monkeypatch):
    # el background task del TestClient corre sincrónico → al volver el POST ya terminó.
    # Monkeypatcheamos la descarga real (red) por una que deja un archivo en la carpeta;
    # el _zip_carpeta REAL lo empaqueta.
    import api.app as modulo

    def fake_descargar_cui(cui, carpeta, *, fecha_inicio=None, fecha_fin=None):
        d = (carpeta / "Valorizaciones" / "2020-09")
        d.mkdir(parents=True, exist_ok=True)
        (d / "doc.pdf").write_bytes(b"%PDF-1.4 demo")
        return {"cui": cui, "obra_id": 111, "nombre": "OBRA DEMO",
                "descargados": 1, "fallidos": 0, "error": None}

    monkeypatch.setattr(modulo, "descargar_cui", fake_descargar_cui)

    r = cliente.post("/api/pivote/descargar-cui",
                     json={"cui": "2418877", "fecha_inicio": "2020-01-01"})
    assert r.status_code == 201
    did = r.json()["descarga_id"]

    r = cliente.get(f"/api/pivote/descargar-cui/{did}")
    assert r.status_code == 200
    body = r.json()
    assert body["estado"] == "listas" and body["listo"] is True
    assert body["obra_id"] == 111 and body["descargados"] == 1

    r = cliente.get(f"/api/pivote/descargar-cui/{did}/zip")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    nombres = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert "indice.txt" in nombres and any(n.endswith("doc.pdf") for n in nombres)


def test_descargar_cui_no_resuelto_marca_error_y_zip_409(cliente, monkeypatch):
    import api.app as modulo
    monkeypatch.setattr(modulo, "descargar_cui",
                        lambda cui, carpeta, **k: {"cui": cui, "obra_id": None,
                                                   "nombre": None, "descargados": 0,
                                                   "fallidos": 0, "error": "no resolvió"})
    did = cliente.post("/api/pivote/descargar-cui",
                       json={"cui": "9999999"}).json()["descarga_id"]
    assert cliente.get(f"/api/pivote/descargar-cui/{did}").json()["estado"] == "error"
    assert cliente.get(f"/api/pivote/descargar-cui/{did}/zip").status_code == 409


# ── Progreso del análisis (barra del panel) ──────────────────────────────────

def test_progreso_endpoint_job_terminado(cliente):
    """Al terminar el pipeline (esqueleto), /progreso muestra 100% y las 8 etapas OK."""
    cid = cliente.post("/api/pivote/concursos",
                       json={"nomenclatura": "CP-PROG/2026"}).json()["concurso_id"]
    job_id = _subir(cliente, cid).json()["job_id"]

    r = cliente.get(f"/api/pivote/jobs/{job_id}/progreso")
    assert r.status_code == 200
    p = r.json()
    assert p["job_id"] == job_id
    assert p["estado"] == "completado"
    assert p["pct"] == 100.0
    assert [e["etapa"] for e in p["etapas"]] == [x.value for x in __import__(
        "schemas.pipeline", fromlist=["Etapa"]).Etapa.orden()]
    assert all(e["estado"] in ("ok", "ok_con_revision", "error_parcial")
               for e in p["etapas"])
    # descargas embebidas + obra en curso (None: ya no corre nada)
    assert "descargas" in p and "obra_actual" in p["descargas"]
    assert p["eta"] is None
    assert p["pendientes_humano"] == 0


def test_progreso_404_si_no_existe(cliente):
    assert cliente.get("/api/pivote/jobs/nope/progreso").status_code == 404


def test_armar_progreso_fusiona_vivo_y_checkpoint(cliente):
    """Etapa ya OK (checkpoint) vs etapa en curso (RegistroProgreso) vs pendiente."""
    import api.app as modulo
    from schemas import pipeline

    job = pipeline.Job(job_id="jx", analisis_id="a")
    job.etapas.append(pipeline.ResultadoEtapa(
        etapa=pipeline.Etapa.INGESTA, estado=pipeline.EstadoEtapa.OK))
    modulo.REGISTRO.reportar("jx", "infoobras", 4, 9, "Verificando la obra 4 de 9")
    try:
        prog = modulo.armar_progreso(job)
        et = {e["etapa"]: e for e in prog["etapas"]}
        assert et["ingesta"]["estado"] == "ok"
        assert et["ingesta"]["texto"] == "Recepción de la propuesta"      # etiqueta estable
        assert et["infoobras"]["estado"] == "en_curso"
        assert et["infoobras"]["item_actual"] == 4
        assert et["infoobras"]["items_total"] == 9
        assert et["infoobras"]["texto"] == "Verificando la obra 4 de 9"   # texto fino en vivo
        assert et["sunat"]["estado"] == "pendiente"
        assert prog["pct"] == round(100 / 8, 1)   # 1 de 8 etapas completa
    finally:
        modulo.REGISTRO.limpiar("jx")


# ── Búsqueda global de profesionales ─────────────────────────────────────────

def test_buscar_profesionales_por_nombre_sin_tildes(cliente):
    cid = cliente.post("/api/pivote/concursos",
                       json={"nomenclatura": "CP-BUSCA/2026"}).json()["concurso_id"]
    job_id = _subir(cliente, cid).json()["job_id"]

    # "profesional uno" matchea "Profesional Uno" (case-insensitive); "unó" prueba tildes
    for q in ("profesional uno", "UNÓ"):
        r = cliente.get("/api/pivote/profesionales", params={"q": q})
        assert r.status_code == 200
        hits = r.json()
        assert any(h["job_id"] == job_id for h in hits), f"sin match para {q!r}"
    h = next(x for x in cliente.get("/api/pivote/profesionales",
                                    params={"q": "profesional uno"}).json()
             if x["job_id"] == job_id)
    assert h["nombre"] == "Profesional Uno"
    assert h["cargo"] == "JEFE DE SUPERVISIÓN"
    assert h["n_prof"] == 1 and h["n_experiencias"] == 1
    assert h["concurso_id"] == cid and h["concurso"] == "CP-BUSCA/2026"
    assert h["postor"] == "POSTOR API DEMO"


def test_buscar_profesionales_por_colegiatura_y_cargo(cliente):
    cid = cliente.post("/api/pivote/concursos",
                       json={"nomenclatura": "CP-BUSCA2/2026"}).json()["concurso_id"]
    job_id = _subir(cliente, cid).json()["job_id"]
    assert any(h["job_id"] == job_id for h in cliente.get(
        "/api/pivote/profesionales", params={"q": "CIP 12345"}).json())
    assert any(h["job_id"] == job_id for h in cliente.get(
        "/api/pivote/profesionales", params={"q": "jefe de supervision"}).json())


def test_buscar_profesionales_query_corta_o_sin_match(cliente):
    assert cliente.get("/api/pivote/profesionales", params={"q": "x"}).json() == []
    assert cliente.get("/api/pivote/profesionales", params={"q": ""}).json() == []
    assert cliente.get("/api/pivote/profesionales",
                       params={"q": "zzz-nadie-zzz"}).json() == []
