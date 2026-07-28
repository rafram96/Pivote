"""
Cosechador de VERDADES para la golden v4 — reconstruye el corpus perdido
(`auditoria_cui_v3.xlsx`, ver `.ai/handoffs/current.md` nota 2026-07-28) desde
la fuente de verdad humana que SÍ existe: las decisiones tomadas en el panel.

Cada ítem de revisión RESUELTO por el evaluador (confirmó la obra o pegó el
CUI correcto) es una fila de verdad humana. Este script recorre TODOS los jobs
de PIVOTE_DATA_DIR y las exporta a `auditoria_cui_v4.xlsx` (hoja
`auditoria_cui`, mismo esquema de columnas que consume `golden_cui.py`):

  col 0 riesgo · 1 job · 2 prof:exp · 3 profesional · 4 cargo · 5 proyecto ·
  6 CUI verdad · 7 nombre obra · 8 CUI citado en cert · 9 RUC emisor ·
  10 fuente de la verdad · 11 nota

Se ejecuta EN EL SERVER (donde viven los jobs reales) o en local; si el
archivo ya existe, AGREGA sin duplicar (dedup por job + prof:exp — la fila
existente gana: una verdad revisada no se pisa con una re-cosecha).

Uso:
  python -m scripts.exportar_verdades_panel                # DATA_DIR → v4
  python -m scripts.exportar_verdades_panel --dry          # solo contar
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

_BACKEND = Path(__file__).resolve().parents[1]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from config import data_dir  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass

HEADER = ["riesgo", "job", "prof_exp", "profesional", "cargo", "proyecto",
          "cui_verdad", "nombre_obra_verdad", "cui_en_cert", "ruc_emisor",
          "fuente_verdad", "nota"]


def _digitos(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def _leer(dj: Path, nombre: str) -> dict:
    ruta = dj / f"{nombre}.json"
    if not ruta.exists():
        return {}
    try:
        return json.loads(ruta.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {}


def cosechar_job(dj: Path) -> list[list]:
    """Filas de verdad de un job: ítems de revisión RESUELTOS cuya resolución
    dejó un CUI en el enriquecimiento (la decisión humana re-ejecutó aguas
    abajo, así que el CUI post-resolución ES la verdad elegida/confirmada)."""
    job = _leer(dj, "job")
    espejo = _leer(dj, "espejo")
    enr = _leer(dj, "enriquecimiento")
    if not job or not espejo:
        return []
    exps = {}
    for p in espejo.get("profesionales") or []:
        for e in p.get("experiencias") or []:
            exps[(p.get("n_prof"), e.get("n"))] = (p, e)
    filas = []
    for it in job.get("items_revision") or []:
        if not it.get("resuelto"):
            continue
        np_, ne = it.get("n_prof"), it.get("n_exp")
        v = (enr or {}).get(f"{np_}:{ne}") or {}
        cui = _digitos(v.get("cui"))
        if not (6 <= len(cui) <= 7):        # sin CUI post-resolución no hay verdad
            continue
        p, e = exps.get((np_, ne), ({}, {}))
        filas.append([
            "REVISION_RESUELTA", dj.name, f"{np_}:{ne}",
            p.get("nombre") or it.get("profesional") or "",
            p.get("cargo") or it.get("cargo") or "",
            e.get("proyecto") or it.get("proyecto") or "",
            cui, str(v.get("obra_nombre") or "")[:120],
            _digitos(e.get("cui")), _digitos(e.get("ruc_emisor")),
            "revision resuelta en el panel (decision humana)",
            f"motivo original: {str(it.get('motivo'))[:80]}",
        ])
    return filas


def main(argv: list[str]) -> int:
    import openpyxl
    base = data_dir()
    destino = base / "auditoria_cui_v4.xlsx"
    nuevas = []
    for dj in sorted(base.iterdir()):
        if dj.is_dir() and (dj / "job.json").exists():
            nuevas.extend(cosechar_job(dj))
    print(f"verdades cosechadas: {len(nuevas)} (de {base})")
    if "--dry" in argv:
        for f in nuevas:
            print("  ", f[1], f[2], "→", f[6], "|", str(f[5])[:60])
        return 0

    if destino.exists():
        wb = openpyxl.load_workbook(destino)
        ws = wb["auditoria_cui"]
        vistas = {(str(r[1]), str(r[2]))
                  for r in ws.iter_rows(min_row=2, values_only=True) if r and r[1]}
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "auditoria_cui"
        ws.append(HEADER)
        vistas = set()

    agregadas = 0
    for f in nuevas:
        if (str(f[1]), str(f[2])) in vistas:
            continue
        ws.append(f)
        agregadas += 1
    wb.save(destino)
    print(f"agregadas {agregadas} filas nuevas → {destino}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
