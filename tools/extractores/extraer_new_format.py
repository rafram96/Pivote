"""
extraer_new_format.py — Extrae la hoja 'CLAUDE' del new_format a un JSON espejo.

Auxiliar de validación (no es parte de la skill). Da un fixture real con los
campos nuevos del formato: experiencia_total_declarada, requisitos por cargo,
cross_checks y notas por profesional, y oferta con ¿inferior?/¿superior?.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parents[2]
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / "fixtures" / "new_format" / "02. Formato de evaluacion COMPLETADO.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else BASE / "fixtures" / "new_format" / "libertador_espejo.json"


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["CLAUDE"]
    MAXR = ws.max_row

    def cell(r, c):
        v = ws.cell(r, c).value
        if isinstance(v, dt.datetime):
            return v.date().isoformat()
        if isinstance(v, str):
            v = v.replace("\n", " ").strip()
            return v or None
        return v

    def txt(r, c):
        v = cell(r, c)
        return v if isinstance(v, str) else (str(v) if v is not None else None)

    def a_starts(r, s):
        v = txt(r, 1)
        return bool(v and v.upper().startswith(s.upper()))

    def find(pred, start=1, end=None):
        for r in range(start, (end or MAXR) + 1):
            if pred(r):
                return r
        return None

    esp = {
        "_meta": {"analisis_id": "libertador-newfmt--2026-06-07T10-20-00",
                  "concurso": "new_format (Consorcio Consultor Libertador)",
                  "postor": txt(2, 3), "version_contrato": "1.1.0",
                  "generado_por": "extraer_new_format.py"},
        "postor": {"detalle": txt(2, 3), "formularios": [], "oferta_economica": {},
                   "experiencia_postor": [], "experiencia_postor_total": {}, "postor_cumple": None},
        "profesionales": [],
        "resumen_evaluacion": {"factores": [], "puntaje_total": None, "nota": None},
    }

    r_p1 = find(lambda r: a_starts(r, "PARTE 1"))
    r_p2 = find(lambda r: a_starts(r, "PARTE 2"))
    r_p5 = find(lambda r: a_starts(r, "PARTE 5"))

    # PARTE 1
    for r in range(r_p1 + 1, r_p2):
        a, b, c, d = txt(r, 1), txt(r, 2), txt(r, 3), txt(r, 4)
        if a and a.strip().lower() == "monto":
            esp["postor"]["oferta_economica"] = {
                "cuantia": cell(r, 2), "limite_inferior": cell(r, 3), "propuesta": cell(r, 4),
                "es_inferior": txt(r, 5), "es_superior": txt(r, 6)}
            continue
        es_header = (b == "DOCUMENTO") or (c or "").upper().startswith("OBSERVACIÓN")
        if (a and (a.upper().startswith("ANEXO") or "REPRESENTACION" in a.upper())) and (b or c) and not es_header:
            esp["postor"]["formularios"].append(
                {"anexo": a or "", "documento": b or "", "observacion": c or "", "folio": d or ""})

    # PARTE 2
    r_h2 = find(lambda r: txt(r, 1) == "No", r_p2, r_p2 + 6)
    for r in range(r_h2 + 1, r_p5):
        a = txt(r, 1)
        if txt(r, 5) and "TOTAL" in (txt(r, 5) or "").upper():
            esp["postor"]["experiencia_postor_total"] = {"acredita": cell(r, 8)}
            continue
        if txt(r, 2) and "CUMPLE" in (txt(r, 2) or "").upper():
            esp["postor"]["postor_cumple"] = txt(r, 8)
            break
        if isinstance(cell(r, 1), (int, float)) and txt(r, 2):
            esp["postor"]["experiencia_postor"].append({
                "n": cell(r, 1), "cliente": txt(r, 2), "contrato": txt(r, 3), "proyecto": txt(r, 4),
                "tipo_acreditacion": txt(r, 5), "monto": cell(r, 6), "pct_objeto": cell(r, 7),
                "le_corresponde": cell(r, 8), "acredita": txt(r, 9), "folio": txt(r, 10),
                "ultimos_20_anios": txt(r, 11), "tipo_solicitado": txt(r, 12), "observaciones": txt(r, 13)})

    # PARTE 3+4 por profesional
    p3_rows = [r for r in range(1, MAXR + 1) if a_starts(r, "PARTE 3:")]
    p3_rows.append(r_p5)
    INFO = {"NOMBRE DEL PROFESIONAL": ("nombre", "folio_nombre"),
            "TÍTULO PROFESIONAL": ("titulo", "folio_titulo"),
            "¿LA PROFESIÓN": ("profesion_valida", None),
            "N° DE COLEGIATURA": ("colegiatura", "fecha_colegiatura"),
            "B. CERTIFICACIONES": ("certificaciones", None),
            "EXPERIENCIA TOTAL DECLARADA": ("experiencia_total_declarada", None)}
    REQ = {"CARGOS VÁLIDOS": "cargos_validos", "TIPO DE EXPERIENCIA": "tipo_experiencia_valida",
           "TIPO DE OBRA": "tipo_obra_valida"}

    for idx in range(len(p3_rows) - 1):
        ini, fin = p3_rows[idx], p3_rows[idx + 1]
        prof = {"n_prof": idx + 1, "experiencias": [], "total": {}, "requisitos": {},
                "cross_checks": [], "notas": []}
        r_p4 = find(lambda r: a_starts(r, "PARTE 4:"), ini, fin) or fin
        # cargo + nombre del título de PARTE 4: "PARTE 4: EXPERIENCIA — CARGO — Nombre"
        t4 = txt(r_p4, 1) or ""
        partes = [s.strip() for s in t4.split("—")]
        prof["cargo"] = partes[1] if len(partes) > 1 else "PROFESIONAL"
        # Parte 3 info
        for r in range(ini, r_p4):
            lab = (txt(r, 3) or "").upper()
            for k, (fi, ff) in INFO.items():
                if lab.startswith(k):
                    prof[fi] = cell(r, 4)
                    if ff:
                        prof[ff] = cell(r, 5)
        # requisitos (B label, C valor)
        r_hdr = find(lambda r: "ENTIDAD" in (txt(r, 2) or "").upper(), r_p4, fin)
        for r in range(r_p4 + 1, r_hdr or fin):
            lb = (txt(r, 2) or "").upper()
            for k, f in REQ.items():
                if k in lb:
                    prof["requisitos"][f] = txt(r, 3)
        # tabla de experiencias
        if r_hdr:
            for r in range(r_hdr + 1, fin):
                k = txt(r, 11)
                if k and k.upper().startswith("TOTAL"):
                    prof["total"] = {"dias": cell(r, 12), "meses": cell(r, 13), "anios": cell(r, 14)}
                    continue
                b = txt(r, 2)
                if b and b.upper() == "NOTA:":
                    if txt(r, 3):
                        prof["notas"].append(txt(r, 3))
                    continue
                if b and txt(r, 3) and not isinstance(cell(r, 1), (int, float)):
                    # fila de cross-check (label en B, valor en C)
                    prof["cross_checks"].append({"label": b, "valor": txt(r, 3)})
                    continue
                if isinstance(cell(r, 1), (int, float)) and txt(r, 2):
                    prof["experiencias"].append({
                        "n": cell(r, 1), "entidad_emisora": txt(r, 2), "proyecto": txt(r, 3),
                        "tipo_documento": txt(r, 4), "nombre_emisor": txt(r, 5), "cargo_emisor": txt(r, 6),
                        "cargo_valido_emitir": txt(r, 7), "fecha_inicial": cell(r, 8), "fecha_final": cell(r, 9),
                        "fecha_emision": cell(r, 10), "folio": txt(r, 11), "dias": cell(r, 12),
                        "meses": cell(r, 13), "anios": cell(r, 14), "anterior_colegiatura": txt(r, 15),
                        "cargo_ocupado": txt(r, 16), "cargo_bases_valido": txt(r, 17),
                        "funciones_similares": txt(r, 18), "cert_antes_culminar": txt(r, 19),
                        "incluye_covid": txt(r, 20), "tipo_obra_valido": txt(r, 21), "observaciones": txt(r, 22)})
        # Reordenar: info del profesional arriba, luego los arrays/objetos.
        ordered = {
            "n_prof": prof["n_prof"],
            "cargo": prof.get("cargo"),
            "nombre": prof.get("nombre"),
            "folio_nombre": prof.get("folio_nombre"),
            "titulo": prof.get("titulo"),
            "folio_titulo": prof.get("folio_titulo"),
            "profesion_valida": prof.get("profesion_valida"),
            "colegiatura": prof.get("colegiatura"),
            "fecha_colegiatura": prof.get("fecha_colegiatura"),
            "certificaciones": prof.get("certificaciones"),
            "experiencia_total_declarada": prof.get("experiencia_total_declarada"),
            "requisitos": prof.get("requisitos", {}),
            "experiencias": prof.get("experiencias", []),
            "total": prof.get("total", {}),
            "cross_checks": prof.get("cross_checks", []),
            "notas": prof.get("notas", []),
        }
        esp["profesionales"].append(ordered)

    # PARTE 5 (new_format: A=factor, B=criterio, C=folio, D=detalle, E=puntaje;
    # el "PUNTAJE TOTAL" lleva el label en col D y el valor en col E).
    for r in range(r_p5 + 1, MAXR + 1):
        a = txt(r, 1)
        dlab = (txt(r, 4) or "")
        if dlab.upper().startswith("PUNTAJE TOTAL") or (a and a.upper().startswith("PUNTAJE")):
            esp["resumen_evaluacion"]["puntaje_total"] = cell(r, 5)
            continue
        if not a:
            continue
        au = a.upper()
        if au.startswith("EL POSTOR"):
            esp["resumen_evaluacion"]["nota"] = a
            continue
        if au == "FACTOR" or au.startswith("CONCLUS"):
            continue
        esp["resumen_evaluacion"]["factores"].append(
            {"factor": a, "criterio": txt(r, 2), "folio": txt(r, 3), "detalle": txt(r, 4), "puntaje": cell(r, 5)})

    OUT.write_text(json.dumps(esp, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    nprof = len(esp["profesionales"])
    nexp = sum(len(p["experiencias"]) for p in esp["profesionales"])
    print(f"OK · {nprof} profesionales · {nexp} experiencias · postor_exp={len(esp['postor']['experiencia_postor'])} · factores={len(esp['resumen_evaluacion']['factores'])}")
    print(f"→ {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
