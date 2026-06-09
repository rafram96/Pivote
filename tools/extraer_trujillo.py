"""
extraer_trujillo.py — Extrae el Excel COMPLETADO de Trujillo a un JSON espejo.

Auxiliar para el round-trip de validación del generador:
    completado.xlsx  →  (este script)  →  trujillo_espejo_full.json
    trujillo_espejo_full.json  →  generar_excel.py  →  _generado_full.xlsx
    comparar _generado_full.xlsx vs completado.xlsx

NO es parte de la skill (la skill produce el JSON con Claude, no extrayendo de un
Excel previo). Esto solo sirve para tener los 13 profesionales reales como fixture
sin transcribir a mano.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parents[1]
SRC = BASE / "fixtures" / "trujillo" / "02. Formato de evaluacion COMPLETADO - Consorcio Salud Trujillo I.xlsx"
OUT = BASE / "fixtures" / "trujillo" / "trujillo_espejo_full.json"


def cell(ws, r, c):
    v = ws.cell(r, c).value
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, str):
        v = v.replace("\n", " ").strip()
        return v or None
    return v


def texto(ws, r, c):
    v = cell(ws, r, c)
    return v if isinstance(v, str) else (str(v) if v is not None else None)


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb.active
    MAXR = ws.max_row

    espejo: dict = {
        "_meta": {
            "analisis_id": "trujillo-cp02-2025--2026-05-30T11-30-00",
            "concurso": "Supervisión Hospital (Consorcio Salud Trujillo I)",
            "postor": "CONSORCIO SALUD TRUJILLO I",
            "version_contrato": "1.0.0",
            "generado_por": "extraer_trujillo.py",
        },
        "postor": {"detalle": texto(ws, 2, 3), "formularios": [], "oferta_economica": {},
                   "experiencia_postor": [], "experiencia_postor_total": {}, "postor_cumple": None},
        "profesionales": [],
        "resumen_evaluacion": {"factores": [], "puntaje_total": None, "nota": None},
    }

    # ── localizar marcadores ──
    def find(pred, start=1, end=None):
        end = end or MAXR
        for r in range(start, end + 1):
            if pred(r):
                return r
        return None

    def a_starts(r, txt):
        v = texto(ws, r, 1)
        return bool(v and v.upper().startswith(txt.upper()))

    r_p1 = find(lambda r: a_starts(r, "PARTE 1"))
    r_p2 = find(lambda r: a_starts(r, "PARTE 2"))
    r_p5 = find(lambda r: a_starts(r, "PARTE 5"))

    # ── PARTE 1: anexos + oferta ──
    for r in range(r_p1 + 1, r_p2):
        anexo, desc, obs, folio = (texto(ws, r, 1), texto(ws, r, 2), texto(ws, r, 3), texto(ws, r, 4))
        if (anexo and anexo.upper().startswith("ANEXO")) or (desc and obs):
            if desc and "Cuantía" not in (desc or ""):
                espejo["postor"]["formularios"].append(
                    {"anexo": anexo or "", "descripcion": desc or "", "observacion": obs or "", "folio": folio or ""})
        # oferta económica: fila con "Monto" en A
        if anexo and anexo.strip().lower() == "monto":
            espejo["postor"]["oferta_economica"] = {
                "cuantia": cell(ws, r, 2), "limite_inferior": cell(ws, r, 3),
                "propuesta": cell(ws, r, 4),
                "detalle": texto(ws, r, 5) or texto(ws, r, 6) or "",
            }

    # ── PARTE 2: experiencia del postor ──
    r_hdr2 = find(lambda r: (texto(ws, r, 1) == "No"), r_p2, r_p2 + 6)
    for r in range(r_hdr2 + 1, r_p5):
        a = texto(ws, r, 1)
        if a and a.upper().startswith("TOTAL"):
            espejo["postor"]["experiencia_postor_total"] = {
                "le_corresponde": cell(ws, r, 8), "acredita": cell(ws, r, 9)}
            continue
        if texto(ws, r, 5) and "CUMPLE" in (texto(ws, r, 5) or "").upper():
            espejo["postor"]["postor_cumple"] = texto(ws, r, 8)
            break
        if isinstance(cell(ws, r, 1), (int, float)) and texto(ws, r, 2):
            espejo["postor"]["experiencia_postor"].append({
                "n": cell(ws, r, 1), "cliente": texto(ws, r, 2), "contrato": texto(ws, r, 3),
                "proyecto": texto(ws, r, 4), "tipo_acreditacion": texto(ws, r, 5),
                "monto": cell(ws, r, 6), "pct_objeto": cell(ws, r, 7),
                "le_corresponde": cell(ws, r, 8), "acredita": cell(ws, r, 9), "folio": texto(ws, r, 10),
                "ultimos_20_anios": texto(ws, r, 11), "tipo_solicitado": texto(ws, r, 12),
                "observaciones": texto(ws, r, 13)})

    # ── PARTE 3+4: bloques por profesional ──
    prof_rows = [r for r in range(1, MAXR + 1) if a_starts(r, "PROFESIONAL ")]
    prof_rows.append(r_p5)  # límite final
    INFO = {"NOMBRE DEL PROFESIONAL": ("nombre", "folio_nombre"),
            "TÍTULO PROFESIONAL": ("titulo", "folio_titulo"),
            "LA PROFESIÓN": ("profesion_valida", None),
            "NO DE COLEGIATURA": ("colegiatura", "folio_colegiatura"),
            "B. CERTIFICACIONES": ("certificaciones", None)}

    for idx in range(len(prof_rows) - 1):
        ini, fin = prof_rows[idx], prof_rows[idx + 1]
        titulo = texto(ws, ini, 1)  # "PROFESIONAL 1: GERENTE DE CONTRATO"
        cargo = titulo.split(":", 1)[1].strip() if ":" in titulo else titulo
        prof: dict = {"n_prof": idx + 1, "cargo": cargo, "experiencias": [], "total": {}}

        r_p4 = find(lambda r: a_starts(r, "PARTE 4"), ini, fin) or fin
        # info general (Parte 3): filas con label en C entre ini y r_p4
        for r in range(ini, r_p4):
            lab = (texto(ws, r, 3) or "").upper()
            for key, (f_info, f_folio) in INFO.items():
                if lab.startswith(key):
                    prof[f_info] = texto(ws, r, 4)
                    if f_folio:
                        prof[f_folio] = texto(ws, r, 5)
        # tabla de experiencias (Parte 4): header con "ENTIDAD" en B
        r_hdr = find(lambda r: "ENTIDAD" in (texto(ws, r, 2) or "").upper(), r_p4, fin)
        if r_hdr:
            for r in range(r_hdr + 1, fin):
                k = texto(ws, r, 11)
                if k and k.upper().startswith("TOTAL"):
                    prof["total"] = {"dias": cell(ws, r, 12), "meses": cell(ws, r, 13), "anios": cell(ws, r, 14)}
                    continue
                a1 = texto(ws, r, 1)
                if a1 and a1.upper().startswith("¿EL PROFESIONAL CUMPLE"):
                    prof["cumple"] = texto(ws, r, 8)
                    continue
                if a1 and a1.upper().startswith("AÑOS ADICIONALES"):
                    prof["anios_adicionales"] = texto(ws, r, 8)
                    continue
                if isinstance(cell(ws, r, 1), (int, float)) and texto(ws, r, 2):
                    prof["experiencias"].append({
                        "n": cell(ws, r, 1), "entidad_emisora": texto(ws, r, 2), "proyecto": texto(ws, r, 3),
                        "tipo_documento": texto(ws, r, 4), "nombre_emisor": texto(ws, r, 5),
                        "cargo_emisor": texto(ws, r, 6), "cargo_valido_emitir": texto(ws, r, 7),
                        "fecha_inicial": cell(ws, r, 8), "fecha_final": cell(ws, r, 9),
                        "fecha_emision": cell(ws, r, 10), "folio": texto(ws, r, 11),
                        "dias": cell(ws, r, 12), "meses": cell(ws, r, 13), "anios": cell(ws, r, 14),
                        "anterior_colegiatura": texto(ws, r, 15), "cargo_ocupado": texto(ws, r, 16),
                        "cargo_bases_valido": texto(ws, r, 17), "funciones_similares": texto(ws, r, 18),
                        "cert_antes_culminar": texto(ws, r, 19), "incluye_covid": texto(ws, r, 20),
                        "tipo_obra_valido": texto(ws, r, 21), "observaciones": texto(ws, r, 22)})
        espejo["profesionales"].append(prof)

    # ── PARTE 5: factores ──
    for r in range(r_p5 + 1, MAXR + 1):
        a = texto(ws, r, 1)
        if not a:
            continue
        if a.upper().startswith("PUNTAJE TÉCNICO"):
            espejo["resumen_evaluacion"]["puntaje_total"] = cell(ws, r, 6)
            continue
        if a.upper().startswith("NOTA"):
            espejo["resumen_evaluacion"]["nota"] = a
            continue
        if a.upper().startswith("FACTOR"):
            continue
        # fila de factor: A=nombre, C=criterio, D=folio, E=detalle, F=puntaje
        if cell(ws, r, 6) is not None or texto(ws, r, 3):
            espejo["resumen_evaluacion"]["factores"].append({
                "factor": a, "criterio": texto(ws, r, 3), "folio": texto(ws, r, 4),
                "detalle": texto(ws, r, 5), "puntaje": cell(ws, r, 6)})

    OUT.write_text(json.dumps(espejo, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    nprof = len(espejo["profesionales"])
    nexp = sum(len(p["experiencias"]) for p in espejo["profesionales"])
    print(f"OK · {nprof} profesionales · {nexp} experiencias · postor_exp={len(espejo['postor']['experiencia_postor'])} · factores={len(espejo['resumen_evaluacion']['factores'])}")
    print(f"→ {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
