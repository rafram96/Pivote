"""
conosce.py — SONDA T-004/T-005: extracción de datos de SEACE por DATOS ABIERTOS
(CONOSCE / OECE), SIN captcha y sin reCAPTCHA.

Descubierto en la sonda del 2026-07-18 como alternativa al buscador con reCAPTCHA
v3 (ver seace.py). El portal CONOSCE publica datasets Excel por año, descargables
por URL directa, actualizados mensualmente (al cierre: 30-jun-2026).

Portal:  https://bi.seace.gob.pe/pentaho/... (páginas dataset.html?pagina=<ds>)
Assets:  https://conosce.osce.gob.pe/buscador/assets/67ae6c4a/reportes/<ds>/<AÑO>/CONOSCE_<DS><AÑO>_0.xlsx

Datasets clave para el pivote (todos se cruzan por `codigoconvocatoria`):
  · convocatorias  → entidad, proceso (nomenclatura), objetocontractual,
                     ubicación, montoreferencial, fecha_convocatoria,
                     fechaintegracionbases, fechapresentacionpropuesta
  · adjudicaciones → proveedor, ruc_proveedor, fecha_buenapro, fecha_consentimiento_bp
  · contratos      → num_contrato, ruc_contratista, monto_contratado_total,
                     **fecha_suscripcion_contrato** (= fecha de firma, T-005),
                     fecha_vigencia_inicial/final, **urlcontrato** (PDF directo)

Hallazgos verificados en vivo desde la laptop (Perú):
  - Contratos 2026: 21.4k filas, 24 cols. urlcontrato baja el PDF del contrato
    SIN captcha (servlet DownloadContratosFileServlet, redirige a prod3).
  - fecha_suscripcion_contrato resuelve T-005 (fecha de firma) sin OCR ni captcha.
  - Cruzando los 3 datasets por codigoconvocatoria se arma la cronología completa
    de un proceso (convocatoria → integración bases → buena pro → suscripción).

Límites conocidos:
  - Actualización mensual (~3 semanas de rezago), NO al día. Suficiente para
    experiencias históricas; procesos de las últimas semanas aún no aparecen.
  - NO trae la URL del PDF de Bases Integradas (sí `fechaintegracionbases`). El
    binario de bases sigue pendiente: probar la API OCDS (record.documents) o el
    servlet de descarga de documentos del procedimiento desde una IP Perú.
  - conosce.osce.gob.pe puede geobloquear IPs fuera de Perú (US → ConnectionError).

Uso:
  python tools/conosce.py contratos 2026           # descarga + resumen de columnas
  python tools/conosce.py convocatorias 2025
  python tools/conosce.py --pdf <urlcontrato> out.pdf   # baja un contrato

Solo lectura. Sonda para T-004/T-005 (aún no es código de producción).
"""
from __future__ import annotations

import sys
from pathlib import Path

import requests

ASSETS = "https://conosce.osce.gob.pe/buscador/assets/67ae6c4a/reportes"
# nombre de carpeta → nombre en el archivo (may/plural)
DATASETS = {
    "contratos": "CONTRATOS",
    "convocatorias": "CONVOCATORIAS",
    "adjudicaciones": "ADJUDICACIONES",
    "pac": "PAC",
    "postor": "POSTOR",            # "Listado de Ofertantes" (consorcios con código interno)
    "proveedores": "PROVEEDORES",  # RNP proveedores individuales
    "consorcios": "CONSORCIO",     # ¡CLAVE!: ruc_consorcio (código) → ruc_miembro + miembro
}
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/144 Safari/537.36")
HEADERS = {"User-Agent": UA, "Referer": "https://bi.seace.gob.pe/"}

OUT = Path(__file__).parent / "_sonda_seace"
OUT.mkdir(parents=True, exist_ok=True)


def url_dataset(ds: str, anio: int) -> str:
    fn = DATASETS[ds]
    return f"{ASSETS}/{ds}/{anio}/CONOSCE_{fn}{anio}_0.xlsx"


def descargar(url: str, dest: Path) -> Path | None:
    try:
        r = requests.get(url, headers=HEADERS, timeout=120, verify=False, stream=True)
        r.raise_for_status()
        dest.write_bytes(r.content)
        return dest
    except Exception as e:  # noqa: BLE001
        print(f"   error: {e}", file=sys.stderr)
        return None


def resumen(xlsx: Path) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    print(f"\n{ws.max_row} filas x {len(hdr)} columnas")
    print("Columnas:")
    for i, h in enumerate(hdr):
        print(f"  {i:2} | {h}")
    print("\nPrimera fila con datos:")
    for row in it:
        if any(v not in (None, "") for v in row):
            for h, v in zip(hdr, row):
                if v not in (None, ""):
                    print(f"   {h}: {str(v)[:72]}")
            break


def main() -> int:
    import urllib3
    urllib3.disable_warnings()

    if len(sys.argv) >= 4 and sys.argv[1] == "--pdf":
        url, dest = sys.argv[2], Path(sys.argv[3])
        print(f"Descargando contrato → {dest} …")
        r = requests.get(url, headers=HEADERS, timeout=120, verify=False)
        dest.write_bytes(r.content)
        ok = r.content[:4] == b"%PDF"
        print(f"  {len(r.content)} bytes | {'PDF válido' if ok else 'NO es PDF'}")
        return 0 if ok else 1

    if len(sys.argv) < 3 or sys.argv[1] not in DATASETS:
        print(__doc__)
        print("Datasets:", ", ".join(DATASETS))
        return 2

    ds, anio = sys.argv[1], int(sys.argv[2])
    url = url_dataset(ds, anio)
    print(f"Descargando {ds} {anio} (sin captcha)…\n  {url}")
    dest = descargar(url, OUT / f"CONOSCE_{ds}_{anio}.xlsx")
    if not dest:
        print("No se pudo descargar (¿geobloqueo fuera de Perú? ¿año sin datos?).")
        return 1
    print(f"  guardado: {dest} ({dest.stat().st_size:,} bytes)")
    resumen(dest)
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass
    raise SystemExit(main())
